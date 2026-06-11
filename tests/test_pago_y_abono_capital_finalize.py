"""Finalize: validaciones PAGO Y ABONO CAPITAL (distribución como control)."""

import asyncio
import io
from datetime import date

import openpyxl
import pytest

from app.application.services.review_schema import (
    ApplicationSubtype,
    CanonicalApplicationType,
    DistribucionCols,
    EstadoPago,
    ExtractRole,
    ReviewSheets,
    TipoAplicacionVisible,
    ValidarPago,
    find_distribucion_pagos_sheet,
)
from app.application.use_cases.payment_validation_finalize import finalize_payment_validation
from tests.test_finalize_validation import (
    MockGraphClient,
    create_review_workbook,
    make_distrib_row,
    set_env_vars,
)


def _run(coro):
    return asyncio.run(coro)


def _policy_fields_pago_y_abono_capital() -> dict[str, object]:
    return {
        DistribucionCols.TIPO_APLICACION_ORIGINAL: TipoAplicacionVisible.PAGO_Y_ABONO_CAPITAL,
        DistribucionCols.TIPO_APLICACION_CANONICA: CanonicalApplicationType.PAGO,
        DistribucionCols.SUBTIPO_APLICACION: ApplicationSubtype.CUOTA_MAS_CAPITAL,
        DistribucionCols.REQUIERE_EXTRACTO: True,
        DistribucionCols.ROL_EXTRACTO: ExtractRole.CIERRE_CUOTA,
        DistribucionCols.CIERRA_CUOTA: True,
        DistribucionCols.ACTUALIZA_IBR: True,
    }


def make_pago_y_abono_capital_row(
    *,
    valor_extracto: float,
    abono_capital: float,
    mora_a_aplicar: float = 0,
    otros: float = 0,
    saldo_por_asignar: float = 0,
    monto_banco: float = 1000,
    estado: str = EstadoPago.NORMAL,
    id_pago: str = "PYAC1",
):
    total = valor_extracto + mora_a_aplicar + abono_capital + otros
    row, hl = make_distrib_row(
        id_pago=id_pago,
        monto_banco=monto_banco,
        valor_int=valor_extracto,
        mora_a_aplicar=mora_a_aplicar,
        abono_capital=abono_capital,
        otros=otros,
        estado=estado,
        validar_pago=ValidarPago.SI,
    )
    vals = dict(zip(DistribucionCols.HEADERS, row))
    vals[DistribucionCols.TOTAL_APLICADO] = total
    vals[DistribucionCols.SALDO_POR_ASIGNAR] = saldo_por_asignar
    vals.update(_policy_fields_pago_y_abono_capital())
    return [vals.get(c, "") for c in DistribucionCols.HEADERS], hl


def test_pago_y_abono_capital_valid_800_200():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_pago_y_abono_capital_row(valor_extracto=800, abono_capital=200)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)]
        )
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = find_distribucion_pagos_sheet(wb)
        cmap = {
            str(ws.cell(1, c).value or "").strip(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(1, c).value
        }
        assert (
            ws.cell(2, cmap[DistribucionCols.SUBTIPO_APLICACION]).value
            == ApplicationSubtype.CUOTA_MAS_CAPITAL
        )

    _run(run_test())


def test_pago_y_abono_capital_rejects_sin_capital():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_pago_y_abono_capital_row(valor_extracto=1000, abono_capital=0)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)]
        )
        with pytest.raises(ValueError, match="pago_y_abono_capital_missing_capital"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())


def test_pago_y_abono_capital_rejects_sin_pago():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_pago_y_abono_capital_row(valor_extracto=0, abono_capital=1000)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)]
        )
        with pytest.raises(ValueError, match="pago_y_abono_capital_missing_parte_cuota"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())


def test_pago_y_abono_capital_rejects_saldo_pendiente():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_pago_y_abono_capital_row(
            valor_extracto=800, abono_capital=150, saldo_por_asignar=50, monto_banco=1000
        )
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)]
        )
        with pytest.raises(ValueError, match="pago_y_abono_capital_saldo_must_be_zero"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())


def test_pago_y_abono_capital_atrasado_cierra_como_pago():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_pago_y_abono_capital_row(
            valor_extracto=700,
            abono_capital=300,
            mora_a_aplicar=0,
            otros=0,
            estado=EstadoPago.ATRASADO,
            monto_banco=1000,
        )
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)]
        )
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = find_distribucion_pagos_sheet(wb)
        cmap = {
            str(ws.cell(1, c).value or "").strip(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(1, c).value
        }
        assert ws.cell(2, cmap[DistribucionCols.ESTADO_PAGO]).value == EstadoPago.ATRASADO
        assert (
            ws.cell(2, cmap[DistribucionCols.TIPO_APLICACION_ORIGINAL]).value
            == TipoAplicacionVisible.PAGO_Y_ABONO_CAPITAL
        )
        assert ws.cell(2, cmap[DistribucionCols.ROL_EXTRACTO]).value == ExtractRole.CIERRE_CUOTA
        assert ws.cell(2, cmap[DistribucionCols.ACTUALIZA_IBR]).value is True
        assert (
            ws.cell(2, cmap[DistribucionCols.SUBTIPO_APLICACION]).value
            != ApplicationSubtype.MORA
        )

    _run(run_test())


def test_distribucion_column_semantics_documented():
    """v2: Mora a aplicar y Abono a capital son columnas independientes; Otros valores aparte."""
    assert DistribucionCols.ABONO_A_CAPITAL == "Abono a capital"
    assert DistribucionCols.MORA_A_APLICAR == "Mora a aplicar"
    assert DistribucionCols.ABONO_A_CAPITAL != DistribucionCols.MORA_A_APLICAR
    assert DistribucionCols.MORA_A_APLICAR in DistribucionCols.HEADERS
    assert DistribucionCols.ABONO_A_CAPITAL in DistribucionCols.HEADERS
    idx_mora = DistribucionCols.HEADERS.index(DistribucionCols.MORA_A_APLICAR)
    idx_cap = DistribucionCols.HEADERS.index(DistribucionCols.ABONO_A_CAPITAL)
    assert idx_mora < idx_cap
