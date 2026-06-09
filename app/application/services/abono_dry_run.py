"""
Dry-run ABONO: preflight documental, cuadre financiero por ID Pago y bloqueo de schedule.

Campo canónico para cuadre: ``PaymentApplicationEvent.valor_pagado_cliente`` (mismo que PAGO).

Cardinalidad PDF → evento:
- ``parse_accounting_text`` produce un único ``PaymentApplicationEvent`` por PDF de asiento.
- ``valor_pagado_cliente`` es el total aplicable del comprobante (no se suman componentes).
- El cuadre del grupo suma ese valor **una vez por ruta normalizada de asiento**; rutas
  repetidas en el mismo grupo no incrementan el total (evita doble conteo si el manifest
  lista el mismo path más de una vez antes de la detección de duplicados).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from app.application.services.accounting_pdf_parser import (
    AccountingParseError,
    PdfTextNotExtractableError,
    PaymentApplicationEvent,
    extract_text_from_pdf,
    parse_accounting_text,
)
from app.application.services.review_schema import (
    DistribucionAbonosCols,
    ReviewSheets,
    TipoAplicacion,
    normalize_credito_digits,
)
from app.application.use_cases.merge_composite_validado_pdfs import normalize_sharepoint_path

# Misma semántica que ``_AMOUNT_TOLERANCE`` en amortization_workbook (0.02).
ABONO_RECONCILIATION_TOLERANCE = Decimal("0.02")

ABONO_MANIFEST_INVALID = "ABONO_MANIFEST_INVALID"
ABONO_MONTO_BANCO_MISSING = "ABONO_MONTO_BANCO_MISSING"
ABONO_FECHA_BANCO_MISSING = "ABONO_FECHA_BANCO_MISSING"
ABONO_CREDIT_ITEMS_MISSING = "ABONO_CREDIT_ITEMS_MISSING"
ABONO_CREDITO_DUPLICADO = "ABONO_CREDITO_DUPLICADO"
ABONO_CREDITO_NO_DECLARADO = "ABONO_CREDITO_NO_DECLARADO"
ABONO_TABLA_AMORTIZACION_MISSING = "ABONO_TABLA_AMORTIZACION_MISSING"
ABONO_ASIENTO_FALTANTE = "ABONO_ASIENTO_FALTANTE"
ABONO_ASIENTO_DUPLICADO = "ABONO_ASIENTO_DUPLICADO"
ABONO_ASIENTO_TOTAL_NOT_FOUND = "ABONO_ASIENTO_TOTAL_NOT_FOUND"
ABONO_ASIENTOS_NO_CUADRAN = "ABONO_ASIENTOS_NO_CUADRAN"
ABONO_SCHEDULE_RULE_NOT_CONFIGURED = "ABONO_SCHEDULE_RULE_NOT_CONFIGURED"


@dataclass
class AbonoCreditItem:
    credito: str
    tipo_aplicacion: str
    ruta_tabla_amortizacion: str
    ruta_unidad_credito: str
    ruta_asientos_contables: str
    asiento_pdf_paths: list[str]
    extracto_pdf_paths: list[str]


@dataclass
class AbonoScheduleContextResult:
    status: str
    reference_date: date | None
    due_date_row: int | None
    application_row: int | None
    ibr_date: date | None
    error_code: str | None = None


@dataclass
class AbonoReconciliationResult:
    reconciliation_status: str
    monto_banco: Decimal | None
    total_asientos: Decimal
    diferencia: Decimal
    tolerancia: Decimal
    credit_amounts: dict[str, Decimal]
    accounting_pdf_amounts: dict[str, Decimal]
    blocking_errors: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class AbonoDryRunGroup:
    bank_code: str
    process_key: str
    id_pago: str
    cliente: str
    monto_banco: Decimal | None
    fecha_banco: date | None
    creditos_seleccionados: list[str]
    credit_items: list[AbonoCreditItem]
    reconciliation_status: str = "PENDING"
    schedule_resolution_status: str = "PENDING"
    group_ready_for_apply: bool = False
    requires_business_rule: bool = False
    blocking_errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def infer_manifest_tipo_aplicacion(output: dict[str, Any]) -> str:
    raw = str(output.get("tipo_aplicacion") or "").strip().upper()
    if raw in (TipoAplicacion.PAGO.value, TipoAplicacion.ABONO.value):
        return raw
    return TipoAplicacion.PAGO.value


def infer_manifest_requiere_extracto(output: dict[str, Any], tipo: str) -> bool:
    if "requiere_extracto" in output:
        return bool(output.get("requiere_extracto"))
    return tipo == TipoAplicacion.PAGO.value


def resolve_abono_schedule_context(
    *,
    group: AbonoDryRunGroup,
    reconciliation: AbonoReconciliationResult,
) -> AbonoScheduleContextResult:
    """
    Regla contractual de fila/IBR para ABONO: pendiente de definición de negocio.
    No infiere fechas ni filas automáticamente.
    """
    if reconciliation.reconciliation_status != "PASSED":
        return AbonoScheduleContextResult(
            status="SKIPPED",
            reference_date=None,
            due_date_row=None,
            application_row=None,
            ibr_date=None,
            error_code=None,
        )
    return AbonoScheduleContextResult(
        status="NOT_CONFIGURED",
        reference_date=None,
        due_date_row=None,
        application_row=None,
        ibr_date=None,
        error_code=ABONO_SCHEDULE_RULE_NOT_CONFIGURED,
    )


def _parse_decimal_amount(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if not text:
        return None
    cleaned = text.replace(".", "").replace(",", ".") if re.search(r",\d{1,2}$", text) else text
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        try:
            return Decimal(re.sub(r"[^\d.\-]", "", text.replace(",", ".")))
        except InvalidOperation:
            return None


def _parse_banco_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _norm_credito_token(raw: str) -> str:
    digits = normalize_credito_digits(raw) or re.sub(r"\D", "", str(raw or ""))
    return digits or str(raw or "").strip()


def _blocking(
    code: str,
    message: str,
    *,
    id_pago: str = "",
    credito: str = "",
    creditos_seleccionados: list[str] | None = None,
    paths: list[str] | None = None,
    next_action: str = "",
) -> dict[str, Any]:
    return {
        "error_code": code,
        "message": message,
        "id_pago": id_pago,
        "credito": credito,
        "creditos_seleccionados": list(creditos_seleccionados or []),
        "paths": list(paths or []),
        "next_action": next_action,
    }


def _credit_items_from_output(output: dict[str, Any]) -> list[AbonoCreditItem]:
    items: list[AbonoCreditItem] = []
    raw_items = output.get("credit_items") or []
    if not isinstance(raw_items, list):
        return items
    for ci in raw_items:
        if not isinstance(ci, dict):
            continue
        paths = [
            str(p).strip().strip("/")
            for p in (ci.get("asiento_pdf_paths") or [])
            if str(p).strip()
        ]
        if not paths:
            single = str(ci.get("asiento_pdf_path") or "").strip().strip("/")
            if single:
                paths = [single]
        extractos = [
            str(p).strip().strip("/")
            for p in (ci.get("extracto_pdf_paths") or [])
            if str(p).strip()
        ]
        items.append(
            AbonoCreditItem(
                credito=str(ci.get("credito") or "").strip(),
                tipo_aplicacion=str(ci.get("tipo_aplicacion") or TipoAplicacion.ABONO.value).strip().upper(),
                ruta_tabla_amortizacion=str(ci.get("ruta_tabla_amortizacion") or "").strip().strip("/"),
                ruta_unidad_credito=str(ci.get("ruta_unidad_credito") or "").strip().strip("/"),
                ruta_asientos_contables=str(ci.get("ruta_asientos_contables") or "").strip().strip("/"),
                asiento_pdf_paths=paths,
                extracto_pdf_paths=extractos,
            )
        )
    return items


def build_abono_group_from_manifest_output(
    output: dict[str, Any],
    *,
    bank_code: str,
    process_key: str,
    abono_hist_index: dict[tuple[str, str], dict[str, Any]] | None = None,
) -> AbonoDryRunGroup:
    id_pago = str(output.get("id_pago") or "").strip()
    cliente = str(output.get("cliente") or "").strip()
    monto = _parse_decimal_amount(output.get("monto_banco"))
    fecha = _parse_banco_date(output.get("fecha_banco"))
    creditos_raw = output.get("creditos_seleccionados") or []
    creditos_sel = [
        _norm_credito_token(str(c))
        for c in creditos_raw
        if str(c).strip()
    ]
    if not creditos_sel:
        legacy = str(output.get("credito") or "").strip()
        if legacy:
            creditos_sel = [_norm_credito_token(p) for p in legacy.split(",") if p.strip()]

    credit_items = _credit_items_from_output(output)
    hist_index = abono_hist_index or {}

    enriched: list[AbonoCreditItem] = []
    for item in credit_items:
        tabla = item.ruta_tabla_amortizacion
        ruta_uc = item.ruta_unidad_credito
        cred_key = _norm_credito_token(item.credito)
        hist = hist_index.get((id_pago, cred_key)) or hist_index.get((id_pago, item.credito))
        if not tabla and hist:
            tabla = str(hist.get("tabla_amortizacion_path") or "").strip().strip("/")
        if not ruta_uc and hist:
            ruta_uc = str(hist.get("ruta_unidad_credito") or "").strip().strip("/")
        enriched.append(
            AbonoCreditItem(
                credito=item.credito or cred_key,
                tipo_aplicacion=item.tipo_aplicacion,
                ruta_tabla_amortizacion=tabla,
                ruta_unidad_credito=ruta_uc,
                ruta_asientos_contables=item.ruta_asientos_contables
                or str((hist or {}).get("ruta_asientos_contables") or "").strip().strip("/"),
                asiento_pdf_paths=list(item.asiento_pdf_paths),
                extracto_pdf_paths=[],
            )
        )

    return AbonoDryRunGroup(
        bank_code=bank_code,
        process_key=process_key,
        id_pago=id_pago,
        cliente=cliente,
        monto_banco=monto,
        fecha_banco=fecha,
        creditos_seleccionados=creditos_sel,
        credit_items=enriched,
    )


def validate_abono_group_structure(group: AbonoDryRunGroup) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    id_pago = group.id_pago
    creditos_sel = list(group.creditos_seleccionados)

    if not id_pago:
        errors.append(
            _blocking(
                ABONO_MANIFEST_INVALID,
                "ID Pago vacío en manifest ABONO",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
            )
        )
    if group.monto_banco is None or group.monto_banco <= 0:
        errors.append(
            _blocking(
                ABONO_MONTO_BANCO_MISSING,
                "Monto bancario inválido o ausente",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
                next_action="Verifique monto_banco en el manifest de Merge.",
            )
        )
    if group.fecha_banco is None:
        errors.append(
            _blocking(
                ABONO_FECHA_BANCO_MISSING,
                "Fecha bancaria inválida o ausente",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
            )
        )
    if not creditos_sel:
        errors.append(
            _blocking(
                ABONO_CREDIT_ITEMS_MISSING,
                "No hay créditos seleccionados en el grupo ABONO",
                id_pago=id_pago,
            )
        )
    if not group.credit_items:
        errors.append(
            _blocking(
                ABONO_CREDIT_ITEMS_MISSING,
                "credit_items vacío en manifest ABONO",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
            )
        )

    seen_credits: set[str] = set()
    declared: set[str] = set()
    for item in group.credit_items:
        cred = _norm_credito_token(item.credito)
        if not cred:
            errors.append(
                _blocking(
                    ABONO_MANIFEST_INVALID,
                    "credit item sin crédito",
                    id_pago=id_pago,
                    creditos_seleccionados=creditos_sel,
                )
            )
            continue
        if cred in seen_credits:
            errors.append(
                _blocking(
                    ABONO_CREDITO_DUPLICADO,
                    f"Crédito {cred} duplicado en credit_items",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )
        seen_credits.add(cred)
        declared.add(cred)

        if not item.ruta_tabla_amortizacion:
            errors.append(
                _blocking(
                    ABONO_TABLA_AMORTIZACION_MISSING,
                    f"Falta ruta de tabla de amortización para crédito {cred}",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                    next_action="Confirme RutaTablaAmortizacion en histórico o manifest.",
                )
            )
        if not item.ruta_asientos_contables:
            errors.append(
                _blocking(
                    ABONO_ASIENTO_FALTANTE,
                    f"Falta RutaAsientosContables para crédito {cred}",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )
        if not item.asiento_pdf_paths:
            errors.append(
                _blocking(
                    ABONO_ASIENTO_FALTANTE,
                    f"Sin asiento_pdf_paths para crédito {cred}",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                    paths=[item.ruta_asientos_contables],
                    next_action="Cargue el PDF del asiento en la carpeta del crédito.",
                )
            )

    for cred in creditos_sel:
        if cred not in declared:
            errors.append(
                _blocking(
                    ABONO_CREDITO_NO_DECLARADO,
                    f"Crédito seleccionado {cred} no tiene credit_item en manifest",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )
    for cred in declared:
        if cred not in set(creditos_sel):
            errors.append(
                _blocking(
                    ABONO_CREDITO_NO_DECLARADO,
                    f"credit_item {cred} no está en creditos_seleccionados",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )

    return errors


def detect_duplicate_asiento_paths(
    groups: list[AbonoDryRunGroup],
) -> dict[str, list[dict[str, str]]]:
    """
    Devuelve paths duplicados -> lista de {id_pago, credito}.
    Incluye duplicados intra-crédito (mismo path repetido en asiento_pdf_paths).
    """
    usage: dict[str, list[dict[str, str]]] = {}
    for group in groups:
        for item in group.credit_items:
            seen_in_credit: set[str] = set()
            for path in item.asiento_pdf_paths:
                norm = normalize_sharepoint_path(path)
                if not norm:
                    continue
                if norm in seen_in_credit:
                    usage.setdefault(norm, []).append(
                        {"id_pago": group.id_pago, "credito": item.credito, "kind": "intra_credit"}
                    )
                seen_in_credit.add(norm)
                usage.setdefault(norm, []).append(
                    {"id_pago": group.id_pago, "credito": item.credito, "kind": "cross_usage"}
                )
    dupes: dict[str, list[dict[str, str]]] = {}
    for path, refs in usage.items():
        if len(refs) > 1:
            dupes[path] = refs
    return dupes


def _canonical_amount_from_event(event: PaymentApplicationEvent) -> Decimal | None:
    """Campo canónico: valor_pagado_cliente (total aplicable del asiento)."""
    vp = event.valor_pagado_cliente
    if vp is None:
        return None
    try:
        amount = Decimal(str(vp))
    except InvalidOperation:
        return None
    if amount <= 0:
        return None
    return amount


async def _download_and_parse_asiento(
    download_fn,
    *,
    id_pago: str,
    cliente: str,
    credito: str,
    asiento_path: str,
) -> tuple[PaymentApplicationEvent | None, dict[str, Any] | None]:
    try:
        pdf_bytes = await download_fn(asiento_path)
        text = extract_text_from_pdf(pdf_bytes)
        event = parse_accounting_text(
            text,
            {
                "id_pago": id_pago,
                "cliente": cliente,
                "credito": credito,
                "asiento_pdf_path": asiento_path,
            },
        )
    except PdfTextNotExtractableError as exc:
        return None, _blocking(
            "PDF_TEXT_NOT_EXTRACTABLE",
            str(exc),
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        )
    except AccountingParseError as exc:
        return None, _blocking(
            getattr(exc, "error_code", None) or "ACCOUNTING_PARSE_FAILED",
            str(exc),
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        )
    except Exception as exc:
        return None, _blocking(
            "ASIENTO_DOWNLOAD_FAILED",
            str(exc)[:500],
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        )

    amount = _canonical_amount_from_event(event)
    if amount is None:
        return None, _blocking(
            ABONO_ASIENTO_TOTAL_NOT_FOUND,
            "No se pudo determinar valor_pagado_cliente del asiento",
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
            next_action="Revise que el PDF del asiento contenga el código 544111100505 con monto.",
        )
    return event, None


async def reconcile_abono_group(
    group: AbonoDryRunGroup,
    download_fn,
    *,
    duplicate_paths: dict[str, list[dict[str, str]]] | None = None,
) -> AbonoReconciliationResult:
    blocking: list[dict[str, Any]] = []
    credit_amounts: dict[str, Decimal] = {}
    pdf_amounts: dict[str, Decimal] = {}
    total = Decimal("0")
    group_seen_pdf_paths: set[str] = set()

    for item in group.credit_items:
        cred = _norm_credito_token(item.credito)
        credit_total = Decimal("0")
        for asiento_path in item.asiento_pdf_paths:
            norm = normalize_sharepoint_path(asiento_path)
            if norm in group_seen_pdf_paths:
                blocking.append(
                    _blocking(
                        ABONO_ASIENTO_DUPLICADO,
                        f"Asiento repetido en el grupo: {norm}",
                        id_pago=group.id_pago,
                        credito=cred,
                        creditos_seleccionados=group.creditos_seleccionados,
                        paths=[norm],
                    )
                )
                continue
            if duplicate_paths and norm in duplicate_paths:
                blocking.append(
                    _blocking(
                        ABONO_ASIENTO_DUPLICADO,
                        f"Asiento reutilizado: {norm}",
                        id_pago=group.id_pago,
                        credito=cred,
                        creditos_seleccionados=group.creditos_seleccionados,
                        paths=[norm],
                        next_action="Use un PDF de asiento distinto por crédito y por ID Pago.",
                    )
                )
                continue
            event, err = await _download_and_parse_asiento(
                download_fn,
                id_pago=group.id_pago,
                cliente=group.cliente,
                credito=cred,
                asiento_path=asiento_path,
            )
            if err:
                blocking.append(err)
                continue
            assert event is not None
            amount = _canonical_amount_from_event(event)
            assert amount is not None
            group_seen_pdf_paths.add(norm)
            credit_total += amount
            pdf_amounts[norm or asiento_path] = amount
            total += amount
        if cred:
            credit_amounts[cred] = credit_total

    monto_banco = group.monto_banco
    if monto_banco is None:
        return AbonoReconciliationResult(
            reconciliation_status="FAILED",
            monto_banco=None,
            total_asientos=total,
            diferencia=Decimal("0"),
            tolerancia=ABONO_RECONCILIATION_TOLERANCE,
            credit_amounts=credit_amounts,
            accounting_pdf_amounts=pdf_amounts,
            blocking_errors=blocking,
        )

    if blocking:
        return AbonoReconciliationResult(
            reconciliation_status="FAILED",
            monto_banco=monto_banco,
            total_asientos=total,
            diferencia=total - monto_banco,
            tolerancia=ABONO_RECONCILIATION_TOLERANCE,
            credit_amounts=credit_amounts,
            accounting_pdf_amounts=pdf_amounts,
            blocking_errors=blocking,
        )

    diferencia = total - monto_banco
    diff_abs = abs(diferencia)
    status = "PASSED" if diff_abs <= ABONO_RECONCILIATION_TOLERANCE else "FAILED"
    if status == "FAILED":
        blocking.append(
            _blocking(
                ABONO_ASIENTOS_NO_CUADRAN,
                "La suma de asientos no cuadra con el monto bancario",
                id_pago=group.id_pago,
                creditos_seleccionados=group.creditos_seleccionados,
                next_action="Revise montos de asientos y monto_banco del manifest.",
            )
        )

    return AbonoReconciliationResult(
        reconciliation_status=status,
        monto_banco=monto_banco,
        total_asientos=total,
        diferencia=diferencia,
        tolerancia=ABONO_RECONCILIATION_TOLERANCE,
        credit_amounts=credit_amounts,
        accounting_pdf_amounts=pdf_amounts,
        blocking_errors=blocking,
    )


def abono_group_result_dict(
    group: AbonoDryRunGroup,
    reconciliation: AbonoReconciliationResult,
    schedule: AbonoScheduleContextResult,
) -> dict[str, Any]:
    return {
        "id_pago": group.id_pago,
        "bank_code": group.bank_code,
        "cliente": group.cliente,
        "monto_banco": float(group.monto_banco) if group.monto_banco is not None else None,
        "fecha_banco": group.fecha_banco.isoformat() if group.fecha_banco else None,
        "total_asientos": float(reconciliation.total_asientos),
        "diferencia": float(reconciliation.diferencia),
        "tolerancia": float(reconciliation.tolerancia),
        "creditos_seleccionados": list(group.creditos_seleccionados),
        "credit_amounts": {k: float(v) for k, v in reconciliation.credit_amounts.items()},
        "accounting_pdf_amounts": {k: float(v) for k, v in reconciliation.accounting_pdf_amounts.items()},
        "reconciliation_status": reconciliation.reconciliation_status,
        "schedule_resolution_status": schedule.status,
        "group_ready_for_apply": group.group_ready_for_apply,
        "requires_business_rule": group.requires_business_rule,
        "blocking_errors": list(group.blocking_errors),
        "warnings": list(group.warnings),
    }


def load_abono_historical_index(hist_bytes: bytes) -> dict[tuple[str, str], dict[str, Any]]:
    """Índice (id_pago, crédito) desde hoja Distribucion_Abonos del histórico."""
    import io

    import openpyxl

    from app.application.services.historical_application_rows import (
        _find_abonos_header_row,
        _find_distribucion_abonos_sheet,
        _get_col_abono,
    )
    from app.application.use_cases.send_validar_extractos_notification import _excel_cell_display

    wb = openpyxl.load_workbook(io.BytesIO(hist_bytes), data_only=True)
    try:
        ws = _find_distribucion_abonos_sheet(wb)
        if ws is None:
            return {}
        h_row, header_map = _find_abonos_header_row(ws)
        col_id = _get_col_abono(header_map, DistribucionAbonosCols.ID_PAGO, "ID Pago")
        col_cred = _get_col_abono(header_map, DistribucionAbonosCols.CREDITO, "Crédito")
        col_cred_norm = _get_col_abono(
            header_map, DistribucionAbonosCols.CREDITO_NORMALIZADO, "CreditoNormalizado"
        )
        col_tabla = _get_col_abono(
            header_map, DistribucionAbonosCols.RUTA_TABLA_AMORTIZACION, "RutaTablaAmortizacion"
        )
        col_uc = _get_col_abono(
            header_map, DistribucionAbonosCols.RUTA_UNIDAD_CREDITO, "RutaUnidadCredito"
        )
        col_ra = _get_col_abono(
            header_map, DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES, "RutaAsientosContables"
        )
        if not col_id or not col_cred:
            return {}
        index: dict[tuple[str, str], dict[str, Any]] = {}
        for r in range(h_row + 1, (ws.max_row or h_row) + 1):
            id_p = _excel_cell_display(ws.cell(r, col_id).value).strip()
            cred_vis = _excel_cell_display(ws.cell(r, col_cred).value).strip()
            if not id_p:
                continue
            cred_norm = ""
            if col_cred_norm:
                cred_norm = _excel_cell_display(ws.cell(r, col_cred_norm).value).strip()
            if not cred_norm:
                cred_norm = _norm_credito_token(cred_vis)
            tabla = ""
            if col_tabla:
                tabla = _excel_cell_display(ws.cell(r, col_tabla).value).strip().strip("/")
            ruta_uc = ""
            if col_uc:
                ruta_uc = _excel_cell_display(ws.cell(r, col_uc).value).strip().strip("/")
            ruta_as = ""
            if col_ra:
                ruta_as = _excel_cell_display(ws.cell(r, col_ra).value).strip().strip("/")
            row_data = {
                "tabla_amortizacion_path": tabla,
                "ruta_unidad_credito": ruta_uc,
                "ruta_asientos_contables": ruta_as,
                "credito_normalizado": cred_norm,
            }
            index[(id_p, cred_norm or cred_vis)] = row_data
            if cred_vis and cred_vis != (cred_norm or cred_vis):
                index.setdefault((id_p, cred_vis), row_data)
        return index
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


async def process_abono_manifest_outputs(
    outputs: list[dict[str, Any]],
    *,
    bank_code: str,
    process_key: str,
    download_fn,
    abono_hist_index: dict[tuple[str, str], dict[str, Any]] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    """
    Procesa outputs ABONO del manifest.

    Devuelve (items observabilidad, abono_group_results, contadores).
    """
    groups = [
        build_abono_group_from_manifest_output(
            out,
            bank_code=bank_code,
            process_key=process_key,
            abono_hist_index=abono_hist_index,
        )
        for out in outputs
        if isinstance(out, dict)
    ]

    counters = {
        "abono_groups_total": len(groups),
        "abono_groups_reconciled": 0,
        "abono_groups_not_reconciled": 0,
        "abono_groups_missing_accounting_pdf": 0,
        "abono_groups_schedule_rule_missing": 0,
        "abono_credit_items_total": sum(len(g.credit_items) for g in groups),
    }

    duplicate_paths = detect_duplicate_asiento_paths(groups)
    all_items: list[dict[str, Any]] = []
    group_results: list[dict[str, Any]] = []

    for group in groups:
        struct_errors = validate_abono_group_structure(group)
        if struct_errors:
            group.blocking_errors.extend(struct_errors)
            if any(e.get("error_code") == ABONO_ASIENTO_FALTANTE for e in struct_errors):
                counters["abono_groups_missing_accounting_pdf"] += 1
            reconciliation = AbonoReconciliationResult(
                reconciliation_status="FAILED",
                monto_banco=group.monto_banco,
                total_asientos=Decimal("0"),
                diferencia=Decimal("0"),
                tolerancia=ABONO_RECONCILIATION_TOLERANCE,
                credit_amounts={},
                accounting_pdf_amounts={},
                blocking_errors=struct_errors,
            )
            schedule = AbonoScheduleContextResult(
                status="SKIPPED",
                reference_date=None,
                due_date_row=None,
                application_row=None,
                ibr_date=None,
            )
            group.reconciliation_status = "FAILED"
            group.schedule_resolution_status = schedule.status
            group.group_ready_for_apply = False
            counters["abono_groups_not_reconciled"] += 1
            group_results.append(abono_group_result_dict(group, reconciliation, schedule))
            all_items.extend(build_abono_observability_items(group, reconciliation, schedule, {}))
            continue

        reconciliation = await reconcile_abono_group(
            group,
            download_fn,
            duplicate_paths=duplicate_paths,
        )
        group.blocking_errors.extend(reconciliation.blocking_errors)
        group.reconciliation_status = reconciliation.reconciliation_status

        parsed_events: dict[str, PaymentApplicationEvent] = {}
        if reconciliation.reconciliation_status == "PASSED" and not reconciliation.blocking_errors:
            counters["abono_groups_reconciled"] += 1
        else:
            counters["abono_groups_not_reconciled"] += 1
            if any(
                e.get("error_code") in (ABONO_ASIENTO_FALTANTE, ABONO_ASIENTO_TOTAL_NOT_FOUND)
                for e in reconciliation.blocking_errors
            ):
                counters["abono_groups_missing_accounting_pdf"] += 1

        schedule = resolve_abono_schedule_context(group=group, reconciliation=reconciliation)
        group.schedule_resolution_status = schedule.status

        if reconciliation.reconciliation_status == "PASSED" and schedule.status == "NOT_CONFIGURED":
            group.requires_business_rule = True
            group.group_ready_for_apply = False
            counters["abono_groups_schedule_rule_missing"] += 1
            group.blocking_errors.append(
                _blocking(
                    ABONO_SCHEDULE_RULE_NOT_CONFIGURED,
                    "El abono y sus asientos contables cuadran, pero todavía no está definida "
                    "la regla contractual para seleccionar la fila y fecha IBR.",
                    id_pago=group.id_pago,
                    creditos_seleccionados=group.creditos_seleccionados,
                    next_action=(
                        "Defina con contabilidad qué cuota, fecha contractual e IBR deben "
                        "utilizarse para los abonos antes de ejecutar Apply."
                    ),
                )
            )
        elif reconciliation.reconciliation_status == "PASSED":
            group.group_ready_for_apply = False

        # Re-parse for observability metadata when reconciliation passed
        if reconciliation.reconciliation_status == "PASSED":
            for item in group.credit_items:
                for asiento_path in item.asiento_pdf_paths:
                    norm = normalize_sharepoint_path(asiento_path)
                    event, _ = await _download_and_parse_asiento(
                        download_fn,
                        id_pago=group.id_pago,
                        cliente=group.cliente,
                        credito=item.credito,
                        asiento_path=asiento_path,
                    )
                    if event:
                        parsed_events[norm or asiento_path] = event

        group_results.append(abono_group_result_dict(group, reconciliation, schedule))
        all_items.extend(
            build_abono_observability_items(group, reconciliation, schedule, parsed_events)
        )

    return all_items, group_results, counters


def build_abono_observability_items(
    group: AbonoDryRunGroup,
    reconciliation: AbonoReconciliationResult,
    schedule: AbonoScheduleContextResult,
    parsed_events: dict[str, PaymentApplicationEvent],
) -> list[dict[str, Any]]:
    """Eventos de observabilidad; ninguno queda listo para Apply."""
    items: list[dict[str, Any]] = []
    schedule_blocked = schedule.status == "NOT_CONFIGURED"
    doc_blocking_errors = [
        e
        for e in (group.blocking_errors or []) + (reconciliation.blocking_errors or [])
        if str(e.get("error_code") or "") != ABONO_SCHEDULE_RULE_NOT_CONFIGURED
    ]
    doc_blocked = reconciliation.reconciliation_status != "PASSED" or bool(doc_blocking_errors)

    for idx, ci in enumerate(group.credit_items, start=1):
        cred = _norm_credito_token(ci.credito)
        for asiento_path in ci.asiento_pdf_paths:
            norm = normalize_sharepoint_path(asiento_path)
            event = parsed_events.get(norm) or parsed_events.get(asiento_path)
            payment_app = None
            if event:
                payment_app = {
                    "valor_pagado_cliente": event.valor_pagado_cliente,
                    "capital": event.capital,
                    "intereses": event.intereses,
                    "mora": event.mora,
                    "retenciones": event.retenciones,
                    "saldos_menores": event.saldos_menores,
                }
            error_code = None
            application_status = "BLOCKED"
            if doc_blocked:
                application_status = "ERROR"
                error_code = (
                    doc_blocking_errors[0].get("error_code")
                    if doc_blocking_errors
                    else ABONO_MANIFEST_INVALID
                )
            elif schedule_blocked:
                application_status = "BLOCKED"
                error_code = ABONO_SCHEDULE_RULE_NOT_CONFIGURED

            items.append(
                {
                    "id_pago": group.id_pago,
                    "cliente": group.cliente,
                    "credito": cred or ci.credito,
                    "tipo_aplicacion": TipoAplicacion.ABONO.value,
                    "requiere_extracto": False,
                    "asiento_pdf_path": asiento_path,
                    "extracto_pdf_path": None,
                    "tabla_amortizacion_path": ci.ruta_tabla_amortizacion or None,
                    "event_index": idx,
                    "payment_application": payment_app,
                    "fecha_limite_pago": None,
                    "due_date_row": None,
                    "application_row": None,
                    "target_row": None,
                    "application_status": application_status,
                    "group_reconciliation_status": reconciliation.reconciliation_status,
                    "schedule_resolution_status": schedule.status,
                    "group_ready_for_apply": False,
                    "requires_business_rule": group.requires_business_rule,
                    "error_code": error_code,
                    "warnings": list(group.warnings),
                }
            )
    return items
