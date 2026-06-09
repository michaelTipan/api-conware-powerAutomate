import asyncio
import io
from datetime import date

import openpyxl
import pytest

from app.application.services.review_schema import (
    AsientosPendientesCols,
    DistribucionAbonosCols,
    DistribucionCols,
    EstadoPago,
    ReviewSheets,
    SUPPORT_NOT_APPLICABLE,
    TipoAplicacion,
    ValidarAbono,
    ValidarPago,
)
from app.application.use_cases.payment_validation_finalize import (
    SECRETARY_FIRST_DATA_ROW,
    SECRETARY_HEADER_ROW,
    SECRETARY_HEADERS,
    SECRETARY_SHEET,
    finalize_payment_validation,
)
from tests.test_finalize_validation import (
    MockGraphClient,
    _run,
    create_review_workbook,
    make_abono_row,
    make_distrib_row,
    set_env_vars,
)


def _skipped_payment_row():
    return make_distrib_row(
        estado=EstadoPago.ATRASADO,
        validar_pago=ValidarPago.NO,
        observacion="Pago no incluido en cierre",
    )[0]


def _finalize_abono_only(client, abono_specs, *, distrib_specs=None):
    r = _skipped_payment_row()
    specs = distrib_specs if distrib_specs is not None else [(r, None)]
    client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
        distrib_specs=specs,
        abono_specs=abono_specs,
    )
    return asyncio.run(
        finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
    )


def _load_uploaded(client, prefix):
    key = next(k for k in client.uploaded_files if prefix in k)
    return key, openpyxl.load_workbook(io.BytesIO(client.uploaded_files[key]))


def test_finalize_legacy_workbook_without_abono_sheet():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)]
        )
        result = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert result["validated_payment_rows"] == 1
        assert result["validated_abono_credit_rows"] == 0
        assert ReviewSheets.DISTRIBUCION_ABONOS not in _load_uploaded(client, "cartera_validada_")[1].sheetnames

    _run(run_test())


def test_finalize_abono_without_selection_fails():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        a1, _ = make_abono_row(id_pago="AB1", validar_abono=ValidarAbono.NO)
        a2, _ = make_abono_row(id_pago="AB1", credito="CRED_B", credito_normalizado="B", validar_abono=ValidarAbono.NO)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(_skipped_payment_row(), None)],
            abono_specs=[(a1, None), (a2, None)],
        )
        with pytest.raises(ValueError, match="abono_without_selected_credit"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())


def test_finalize_abono_single_credit_success():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        a1, tab = make_abono_row(id_pago="AB1", validar_abono=ValidarAbono.SI)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(_skipped_payment_row(), None)],
            abono_specs=[(a1, tab)],
        )
        result = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert result["validated_abono_groups"] == 1
        assert result["validated_abono_credit_rows"] == 1
        assert result["validated_rows"] == 1
        assert result["support_abono_rows"] == 1
        assert result.get("user_message")

        _, wb_hist = _load_uploaded(client, "cartera_validada_")
        assert ReviewSheets.DISTRIBUCION_ABONOS in wb_hist.sheetnames

        _, wb_sec = _load_uploaded(client, "soporte_asientos_contables_")
        ws = wb_sec[SECRETARY_SHEET]
        headers = [ws.cell(SECRETARY_HEADER_ROW, c).value for c in range(1, len(SECRETARY_HEADERS) + 1)]
        assert headers == SECRETARY_HEADERS
        assert len(headers) == 13
        assert (
            ws.cell(SECRETARY_FIRST_DATA_ROW, SECRETARY_HEADERS.index("Tipo Aplicación") + 1).value
            == TipoAplicacion.ABONO.value
        )
        assert (
            ws.cell(SECRETARY_FIRST_DATA_ROW, SECRETARY_HEADERS.index("Fecha límite") + 1).value
            == SUPPORT_NOT_APPLICABLE
        )
        assert (
            ws.cell(SECRETARY_FIRST_DATA_ROW, SECRETARY_HEADERS.index("Total validado") + 1).value
            == SUPPORT_NOT_APPLICABLE
        )
        assert (
            ws.cell(SECRETARY_FIRST_DATA_ROW, SECRETARY_HEADERS.index("Link extracto") + 1).value
            == SUPPORT_NOT_APPLICABLE
        )

    _run(run_test())


def test_finalize_abono_multiple_credits_two_support_rows():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        a1, t1 = make_abono_row(id_pago="AB2", credito="CRED", credito_normalizado="CRED", validar_abono=ValidarAbono.SI)
        a2, t2 = make_abono_row(
            id_pago="AB2",
            credito="CRED_B",
            credito_normalizado="B",
            ruta_unidad_credito="clientes/CLI/CRED_B",
            ruta_tabla_amortizacion="clientes/CLI/CRED_B/tabla.xlsx",
            validar_abono=ValidarAbono.SI,
        )
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(_skipped_payment_row(), None)],
            abono_specs=[(a1, t1), (a2, t2)],
        )
        result = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert result["validated_abono_credit_rows"] == 2
        assert result["support_abono_rows"] == 2

        _, wb_sec = _load_uploaded(client, "soporte_asientos_contables_")
        ws = wb_sec[SECRETARY_SHEET]
        rows = []
        for r in range(SECRETARY_FIRST_DATA_ROW, ws.max_row + 1):
            tipo = ws.cell(r, AsientosPendientesCols.HEADERS.index(AsientosPendientesCols.TIPO_APLICACION) + 1).value
            if tipo == TipoAplicacion.ABONO.value:
                rows.append(r)
        assert len(rows) == 2
        montos = [
            ws.cell(r, AsientosPendientesCols.HEADERS.index(AsientosPendientesCols.MONTO_BANCO) + 1).value
            for r in rows
        ]
        assert montos[0] == montos[1] == 100000

    _run(run_test())


def test_finalize_abono_duplicate_credit_fails():
    set_env_vars()
    client = MockGraphClient()
    a1, _ = make_abono_row(id_pago="AB3", credito="CRED", credito_normalizado="CRED", validar_abono=ValidarAbono.SI)
    a2, _ = make_abono_row(id_pago="AB3", credito="CRED", credito_normalizado="CRED", validar_abono=ValidarAbono.SI)
    client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
        distrib_specs=[(_skipped_payment_row(), None)],
        abono_specs=[(a1, None), (a2, None)],
    )
    with pytest.raises(ValueError, match="abono_duplicate_selected_credit"):
        asyncio.run(finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10)))


def test_finalize_abono_group_inconsistent_monto_fails():
    set_env_vars()
    client = MockGraphClient()
    a1, _ = make_abono_row(id_pago="AB4", monto_banco=100000, validar_abono=ValidarAbono.SI)
    a2, _ = make_abono_row(id_pago="AB4", credito="CRED_B", credito_normalizado="B", monto_banco=50000, validar_abono=ValidarAbono.SI)
    client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
        distrib_specs=[(_skipped_payment_row(), None)],
        abono_specs=[(a1, None), (a2, None)],
    )
    with pytest.raises(ValueError, match="abono_group_inconsistent"):
        asyncio.run(finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10)))


def test_finalize_mixed_pago_and_abono():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, tab = make_distrib_row()
        a1, at = make_abono_row(id_pago="AB5", validar_abono=ValidarAbono.SI)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, tab)],
            abono_specs=[(a1, at)],
        )
        result = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert result["validated_payment_rows"] == 1
        assert result["validated_abono_credit_rows"] == 1
        assert result["validated_rows"] == 2
        assert result["payments_and_abonos_supported"] is True

        _, wb_sec = _load_uploaded(client, "soporte_asientos_contables_")
        ws = wb_sec[SECRETARY_SHEET]
        tipos = []
        for ridx in range(SECRETARY_FIRST_DATA_ROW, ws.max_row + 1):
            tipos.append(ws.cell(ridx, 1).value)
        assert TipoAplicacion.PAGO.value in tipos
        assert TipoAplicacion.ABONO.value in tipos

    _run(run_test())


def test_finalize_abono_invalid_validar_abono_fails():
    set_env_vars()
    client = MockGraphClient()
    a1, _ = make_abono_row(validar_abono="FOO")
    client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
        distrib_specs=[(_skipped_payment_row(), None)],
        abono_specs=[(a1, None)],
    )
    with pytest.raises(ValueError, match="invalid_validar_abono"):
        asyncio.run(finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10)))

