import asyncio
import io
from datetime import date
from unittest import mock

import openpyxl
import pytest

from app.application.services.review_schema import (
    DistribucionAbonosCols,
    DistribucionCols,
    ReviewSheets,
    TipoAplicacion,
    ValidarAbono,
    normalize_tipo_aplicacion,
)
from app.application.use_cases.payment_validation_generate import generate_payment_validation
from tests.test_generate_validation import (
    MockGraphClient,
    STANDARD_BANK_HEADERS,
    create_bank_excel,
    create_excel,
    create_pdf_bytes,
    load_generated_workbook,
    make_pdf_extractor_mock,
    set_env_vars,
    setup_client_structure,
    sheet_to_dicts,
)


def _fake_fecha_limite_from_marker_bytes():
    from datetime import date as date_cls

    def _fake(pdf_bytes: bytes):
        text = pdf_bytes.decode(errors="ignore")
        if "FECHA_LIMITE=" in text:
            raw = text.split("FECHA_LIMITE=", 1)[1].split(";", 1)[0].strip()
            return date_cls.fromisoformat(raw)
        return date_cls(2025, 12, 23)

    return _fake


def _run_with_pdf_mock(client):
    extractor = make_pdf_extractor_mock(client)
    return mock.patch(
        "app.application.use_cases.payment_validation_generate.extract_total_a_pagar_from_pdf",
        side_effect=extractor,
    ), mock.patch(
        "app.application.use_cases.payment_validation_generate.extract_fecha_limite_pago_from_pdf",
        side_effect=_fake_fecha_limite_from_marker_bytes(),
    )


@pytest.mark.parametrize(
    "header",
    ["Tipo Aplicación", "Tipo Aplicacion", "tipo aplicacion", "tipoaplicacion"],
)
def test_bank_header_tipo_aplicacion_normalized_variants(header):
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        headers = ["Fecha", "Crédito", "Concepto", header, "Transacción"]
        client.downloaded_files["banco.xlsx"] = create_excel(
            headers,
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "tx-1"]],
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            result = await generate_payment_validation(client, date(2026, 5, 26))
        assert result["pagos_detectados"] == 1
        wb = load_generated_workbook(client)
        assert ReviewSheets.DISTRIBUCION_PAGOS in wb.sheetnames
        assert ReviewSheets.DISTRIBUCION not in wb.sheetnames

    asyncio.run(_run())


def test_bank_header_missing_tipo_aplicacion_fails_before_workbook():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_excel(
            ["Fecha", "Crédito", "Concepto", "Transacción"],
            [[date(2025, 12, 23), 1000, "GEOEXCON", "tx"]],
        )
        with pytest.raises(ValueError, match="tipo_aplicacion_column_missing"):
            with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
                await generate_payment_validation(client, date(2026, 5, 26))
        assert not client.uploaded_files

    asyncio.run(_run())


def test_bank_header_duplicate_tipo_aplicacion_fails():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        client.downloaded_files["banco.xlsx"] = create_excel(
            ["Fecha", "Crédito", "Concepto", "Tipo Aplicación", "Tipo Aplicacion", "Transacción"],
            [[date(2025, 12, 23), 1000, "GEOEXCON", "PAGO", "PAGO", "tx"]],
        )
        with pytest.raises(ValueError, match="tipo_aplicacion_column_duplicate"):
            await generate_payment_validation(client, date(2026, 5, 26))
        assert not client.uploaded_files

    asyncio.run(_run())


def test_transaccion_column_can_be_in_position_e():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        headers = ["Fecha", "Monto", "Concepto", "Tipo Aplicación", "Transacción"]
        client.downloaded_files["banco.xlsx"] = create_excel(
            headers,
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "REF-99"]],
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            result = await generate_payment_validation(client, date(2026, 5, 26))
        wb = load_generated_workbook(client)
        casos = sheet_to_dicts(wb[ReviewSheets.CASOS_PAGO])
        assert len(casos) >= 1
        assert result["pagos_detectados"] == 1

    asyncio.run(_run())


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("PAGO", TipoAplicacion.PAGO),
        ("pago", TipoAplicacion.PAGO),
        ("PAGO Y ABONO CAPITAL", TipoAplicacion.PAGO),
        ("ABONO CAPITAL", TipoAplicacion.ABONO),
        ("ABONO MORA", TipoAplicacion.ABONO),
        ("abono", TipoAplicacion.ABONO),
    ],
)
def test_normalize_tipo_aplicacion_values(raw, expected):
    assert normalize_tipo_aplicacion(raw) == expected


@pytest.mark.parametrize("raw", ["", "FOO", "P", "Pago normal", "SI"])
def test_normalize_tipo_aplicacion_rejects_invalid(raw):
    with pytest.raises(ValueError, match="tipo_aplicacion"):
        normalize_tipo_aplicacion(raw)


def test_row_empty_tipo_fails_fast():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "", "tx"]],
        )
        with pytest.raises(ValueError, match="tipo_aplicacion_required"):
            with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
                await generate_payment_validation(client, date(2026, 5, 26))
        assert not client.uploaded_files

    asyncio.run(_run())


def test_row_generic_abono_fails_fast():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "ABONO", "tx"]],
        )
        with pytest.raises(ValueError, match="generic_abono_not_supported"):
            with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
                await generate_payment_validation(client, date(2026, 5, 26))
        assert not client.uploaded_files

    asyncio.run(_run())


def test_row_invalid_tipo_fails_fast():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "FOO", "tx"]],
        )
        with pytest.raises(ValueError, match="tipo_aplicacion_invalid"):
            with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
                await generate_payment_validation(client, date(2026, 5, 26))
        assert not client.uploaded_files

    asyncio.run(_run())


def test_blank_and_format_rows_do_not_trigger_tipo_errors():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(STANDARD_BANK_HEADERS)
        ws.append([date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "tx"])
        ws.append([None, None, None, None, None])
        ws.append(["", "", "", "", ""])
        ws.append(["Fecha", "Crédito", "Concepto", "Tipo Aplicación", "Transacción"])
        ws.append(["TOTAL", None, None, None, None])
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["banco.xlsx"] = buf.getvalue()
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            result = await generate_payment_validation(client, date(2026, 5, 26))
        assert result["summary"]["transacciones_banco"] == 1

    asyncio.run(_run())


def test_abono_creates_distribucion_abonos_sheet():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 500000, "GEOEXCON", "ABONO CAPITAL", "tx-ab"]],
            default_tipo="ABONO CAPITAL",
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            result = await generate_payment_validation(client, date(2026, 5, 26))
        wb = load_generated_workbook(client)
        assert ReviewSheets.DISTRIBUCION_ABONOS in wb.sheetnames
        assert ReviewSheets.DISTRIBUCION_PAGOS in wb.sheetnames
        assert ReviewSheets.DISTRIBUCION not in wb.sheetnames
        abono_rows = sheet_to_dicts(wb[ReviewSheets.DISTRIBUCION_ABONOS])
        assert len(abono_rows) >= 2
        ids = {r[DistribucionAbonosCols.ID_PAGO] for r in abono_rows}
        assert len(ids) == 1
        assert all(r[DistribucionAbonosCols.VALIDAR_ABONO] == ValidarAbono.NO for r in abono_rows)
        assert DistribucionAbonosCols.LINK_EXTRACTO in DistribucionAbonosCols.HEADERS
        assert DistribucionCols.ESTADO_PAGO not in DistribucionAbonosCols.HEADERS
        assert result["abonos_detectados"] == 1
        assert result["distribution_abonos_sheet"] == ReviewSheets.DISTRIBUCION_ABONOS

    asyncio.run(_run())


def test_mixed_pago_and_abono_workbook():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_excel(
            STANDARD_BANK_HEADERS,
            [
                [date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "tx-p"],
                [date(2025, 12, 24), 100000, "GEOEXCON", "ABONO CAPITAL", "tx-a"],
            ],
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            result = await generate_payment_validation(client, date(2026, 5, 26))
        wb = load_generated_workbook(client)
        pago_rows = sheet_to_dicts(wb[ReviewSheets.DISTRIBUCION_PAGOS])
        abono_rows = sheet_to_dicts(wb[ReviewSheets.DISTRIBUCION_ABONOS])
        assert len(pago_rows) >= 1
        assert len(abono_rows) >= 1
        assert result["pagos_detectados"] == 1
        assert result["abonos_detectados"] == 1
        assert result["summary"]["filas_distribucion_pagos"] == len(pago_rows)
        assert result["summary"]["filas_distribucion_abonos"] == len(abono_rows)
        assert result.get("user_message")

    asyncio.run(_run())


def test_distribucion_sheet_keeps_23_columns():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "tx"]],
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            await generate_payment_validation(client, date(2026, 5, 26))
        wb = load_generated_workbook(client)
        ws = wb[ReviewSheets.DISTRIBUCION_PAGOS]
        hr = 3
        headers = [ws.cell(hr, c).value for c in range(1, ws.max_column + 1)]
        assert len([h for h in headers if h]) == len(DistribucionCols.HEADERS)

    asyncio.run(_run())


def test_abono_sheet_always_created_even_without_abono_rows():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "tx"]],
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            await generate_payment_validation(client, date(2026, 5, 26))
        wb = load_generated_workbook(client)
        assert ReviewSheets.DISTRIBUCION_ABONOS in wb.sheetnames
        abono_rows = sheet_to_dicts(wb[ReviewSheets.DISTRIBUCION_ABONOS])
        assert abono_rows == []

    asyncio.run(_run())
