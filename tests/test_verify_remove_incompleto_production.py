"""Verificación preproducción: sin acceso productivo a pagos_incompletos y schema Generate."""

from __future__ import annotations

import asyncio
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from openpyxl import load_workbook

from app.application.services.historical_application_rows import read_validated_payment_rows
from app.application.services.review_schema import (
    DistribucionAbonosCols,
    DistribucionCols,
    EstadoPago,
    ReviewSheets,
    ValidarPago,
)
from app.application.use_cases.setup_payment_followup_workbooks import (
    ADELANTADOS_COLUMNS,
    SHEET_HISTORICO,
    SHEET_PENDIENTES,
)
from tests.test_finalize_validation import (
    MockGraphClient,
    create_review_workbook,
    make_distrib_row,
    set_env_vars,
)
from tests.test_generate_validation import run_generate
from tests.test_payment_followup_finalize import (
    CONTROL_FOLDER,
    FILENAME_ADELANTADOS,
    PATH_AD,
    MockGraphFollowup,
    _build_followup_workbook_bytes,
    _seed_workbook,
)
from app.application.services.payment_followup_finalize import register_payment_followups_after_finalize
from app.application.use_cases.payment_validation_finalize import finalize_payment_validation
FORBIDDEN_GRAPH_SUBSTRINGS = (
    "pagos_incompletos",
    "pagos_incompletos.xlsx",
    "GRAPH_FOLLOWUP_PAGOS_INCOMPLETOS_PATH",
    "GRAPH_AUDIT_PAGOS_INCOMPLETOS_PATH",
)


def _assert_no_forbidden_paths(paths: list[str]) -> None:
    for p in paths:
        low = p.lower()
        for bad in FORBIDDEN_GRAPH_SUBSTRINGS:
            assert bad.lower() not in low, f"Ruta prohibida detectada: {p!r} (match {bad!r})"


def test_generate_workbook_listas_and_validation_exclude_incompleto():
    _, _, workbook = run_generate(
        [[__import__("datetime").datetime(2025, 12, 23), 25443565, "GEOEXCON", "trx-geo"]],
        date(2025, 12, 23),
    )
    ws_lists = workbook[ReviewSheets.LISTAS]
    dist_statuses = [ws_lists[f"C{idx}"].value for idx in range(1, 10) if ws_lists[f"C{idx}"].value]
    assert "INCOMPLETO" not in dist_statuses
    assert dist_statuses == list(EstadoPago.OPTIONS_ORDERED)

    ws_dist = workbook[ReviewSheets.DISTRIBUCION_PAGOS]
    formulas = [dv.formula1 for dv in ws_dist.data_validations.dataValidation]
    assert any("=_Listas!$C$1:$C$4" in f for f in formulas)
    assert not any("INCOMPLETO" in (f or "") for f in formulas)

    from app.application.use_cases.payment_validation_generate import _ESTADO_PAGO_FILLS

    assert "INCOMPLETO" not in {str(k) for k in _ESTADO_PAGO_FILLS}

    if ReviewSheets.DISTRIBUCION_ABONOS in workbook.sheetnames:
        ws_ab = workbook[ReviewSheets.DISTRIBUCION_ABONOS]
        headers = [str(ws_ab.cell(1, c).value or "") for c in range(1, ws_ab.max_column + 1)]
        assert DistribucionCols.ESTADO_PAGO not in headers
        assert "Estado Pago" not in headers


def test_historical_incompleto_text_still_readable_for_notify_merge():
    """Históricos cerrados con texto INCOMPLETO no se rechazan en lectura Notify/Merge."""
    r, _ = make_distrib_row(
        id_pago="HIST1",
        validar_pago=ValidarPago.SI,
        estado="INCOMPLETO",
        ruta_pdf_internal="ext/hist1.pdf",
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.DISTRIBUCION
    ws.append(DistribucionCols.HEADERS)
    ws.append(r)
    bio = BytesIO()
    wb.save(bio)
    loaded = load_workbook(BytesIO(bio.getvalue()), data_only=True)
    rows = read_validated_payment_rows(loaded, legacy_estado_token="VALIDAR")
    assert len(rows) == 1
    assert rows[0]["id_pago"] == "HIST1"


def test_finalize_adelantado_only_touches_pagos_adelantados_followup():
    g = MockGraphFollowup()
    _seed_workbook(g)
    asyncio.run(
        register_payment_followups_after_finalize(
            g,
            "site-1",
            "drive-1",
            process_date=date(2026, 5, 10),
            distributions=[
                {
                    DistribucionCols.ID_PAGO: "P1",
                    DistribucionCols.CLIENTE: "C",
                    DistribucionCols.CREDITO: "CR1",
                    DistribucionCols.FECHA_BANCO: "2026-05-10",
                    DistribucionCols.FECHA_LIMITE: "2026-05-15",
                    DistribucionCols.MONTO_BANCO: 100.0,
                    DistribucionCols.VALOR_EXTRACTO: 100.0,
                    DistribucionCols.APLICAR_A_EXTRACTO: 100.0,
                    DistribucionCols.ABONO_A_CAPITAL: 0.0,
                    DistribucionCols.OTROS_VALORES: 0.0,
                    DistribucionCols.TOTAL_APLICADO: 100.0,
                    DistribucionCols.SALDO_POR_ASIGNAR: 0.0,
                    DistribucionCols.ESTADO_PAGO: EstadoPago.ADELANTADO,
                    DistribucionCols.VALIDAR_PAGO: ValidarPago.SI,
                    DistribucionCols.OBSERVACION: "",
                    DistribucionCols.LINK_TABLA: "/t",
                    DistribucionCols.RUTA_UNIDAD_CREDITO: "/c",
                    DistribucionCols.RUTA: "/e",
                    DistribucionCols.RUTA_ASIENTOS_CONTABLES: "/a",
                }
            ],
            historical_relative_path="hist/cartera.xlsx",
        )
    )
    _assert_no_forbidden_paths(g.get_bytes_calls)
    assert PATH_AD in g.files
    assert ADELANTADOS_COLUMNS[0] == "ID Pago"
    wb = openpyxl.load_workbook(BytesIO(g.files[PATH_AD]))
    assert SHEET_PENDIENTES in wb.sheetnames
    assert SHEET_HISTORICO in wb.sheetnames


def test_finalize_success_graph_paths_exclude_incompletos():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(estado=EstadoPago.NORMAL)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        path_ad = f"{CONTROL_FOLDER}/{FILENAME_ADELANTADOS}"
        client.downloaded_files[path_ad] = _build_followup_workbook_bytes(columns=ADELANTADOS_COLUMNS)
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        _assert_no_forbidden_paths(client.requested_file_paths)
        _assert_no_forbidden_paths([p for p, _ in client.put_calls])

    asyncio.run(run_test())


def test_pagos_adelantados_columns_unchanged():
    assert ADELANTADOS_COLUMNS == (
        "ID Pago",
        "Cliente",
        "Crédito",
        "FechaPago",
        "FechaLimitePago",
        "FechaIBRRequerida",
        "MontoBanco",
        "ValorExtracto",
        "AplicarAExtracto",
        "MoraAAplicar",
        "OtrosValores",
        "TotalAplicado",
        "SaldoPorAsignar",
        "TablaAmortizacionPath",
        "RutaUnidadCredito",
        "RutaExtracto",
        "AsientoPdfPath",
        "HistoricalFilePath",
        "FilaAplicacionPago",
        "FilaIBR",
        "EstadoAplicacionPago",
        "EstadoIBR",
        "EstadoFinal",
        "IBRUsado",
        "FechaCierreIBR",
        "Observacion",
        "CreatedAt",
        "UpdatedAt",
    )
