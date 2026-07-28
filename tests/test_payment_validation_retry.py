import asyncio
import io
from datetime import date
from unittest import mock
import pytest
import openpyxl

from app.application.use_cases.payment_validation_finalize import finalize_payment_validation
from app.application.use_cases.payment_validation_generate import generate_payment_validation
from tests.test_finalize_validation import MockGraphClient, set_env_vars
from app.application.config.payment_validation_settings import resolve_bank_control_file_path
from app.application.use_cases.setup_merge_control_workbook import _build_process_control_workbook_bytes

def create_finalize_excel_with_errors(bank_code: str, process_date: date, error_rows: int = 1) -> bytes:
    from app.application.services.review_schema import ReviewSheets, ControlCols, DistribucionCols
    wb = openpyxl.Workbook()
    
    # Setup Control sheet
    ws_ctrl = wb.active
    ws_ctrl.title = ReviewSheets.CONTROL
    ws_ctrl.append([ControlCols.CAMPO, ControlCols.VALOR])
    ws_ctrl.append([ControlCols.ROW_ID_PROCESO, f"payment-validation|{bank_code}|{process_date.isoformat()}"])
    ws_ctrl.append([ControlCols.ROW_PROCESAR, "SI"])
    ws_ctrl.append([ControlCols.ROW_ESTADO, "EN_REVISION"])
    
    # Setup Distribucion_Pagos sheet
    ws_dist = wb.create_sheet(ReviewSheets.DISTRIBUCION_PAGOS)
    from app.application.services.review_schema import DistribucionCols
    headers = list(DistribucionCols.HEADERS)
    ws_dist.append(headers)
    
    # Create a valid distrib row by mapping headers
    row_data = {h: "" for h in headers}
    row_data[DistribucionCols.ID_PAGO] = "1001"
    row_data[DistribucionCols.SUBTIPO_APLICACION] = "PAGO"
    row_data[DistribucionCols.ESTADO_PAGO] = "NORMAL"
    row_data[DistribucionCols.VALIDAR_PAGO] = "NO"
    row_data[DistribucionCols.OBSERVACION] = "Testing validation"
    row_data[DistribucionCols.CLIENTE] = "CLIENT"
    row_data[DistribucionCols.APLICAR_A_EXTRACTO] = "0"
    row_data[DistribucionCols.MORA_A_APLICAR] = "0"
    row_data[DistribucionCols.ABONO_A_CAPITAL] = "0"
    row_data[DistribucionCols.OTROS_VALORES] = "0"
    ws_dist.append([row_data[h] for h in headers])
    
    # Setup Distribucion_Abonos sheet
    ws_abonos = wb.create_sheet(ReviewSheets.DISTRIBUCION_ABONOS)
    from app.application.services.review_schema import DistribucionAbonosCols
    headers_abonos = list(DistribucionAbonosCols.HEADERS)
    ws_abonos.append(headers_abonos)
    
    # Setup Errores sheet
    ws_err = wb.create_sheet(ReviewSheets.ERRORES)
    ws_err.append(["Error Code", "Message"])
    for i in range(error_rows):
        ws_err.append(["customer_not_found", f"Error row {i+1}"])
        
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def override_control_state(bank_code: str, estado: str) -> bytes:
        from openpyxl import load_workbook
        import io
        base_bytes = _build_process_control_workbook_bytes(bank_code, "Banco")
        wb = load_workbook(io.BytesIO(base_bytes))
        ws = wb.active
        # Columns in control file: Id, ProcessKey, Banco, EstadoProceso, ValidationFilePath, HistoricalFilePath, SecretaryFilePath...
        # Let's find EstadoProceso column
        headers = {c.value: c.column for c in ws[1]}
        if "EstadoProceso" in headers:
            ws.cell(2, headers["EstadoProceso"], value=estado)
        if "ProcessKey" in headers:
            ws.cell(2, headers["ProcessKey"], value=f"payment-validation|{bank_code}|2026-06-25")
        if "ValidationFilePath" in headers:
            ws.cell(2, headers["ValidationFilePath"], value=f"01 REVISION/revision_banco.xlsx")
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

def test_finalize_bloquea_por_estado_de_correccion():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        
        ctrl_path = resolve_bank_control_file_path("banco_bogota")
        client.downloaded_files[ctrl_path] = override_control_state("banco_bogota", "REVISION_REQUIERE_CORRECCION")
        client.downloaded_files["revision/revision_banco.xlsx"] = create_finalize_excel_with_errors("banco_bogota", date(2026, 6, 25), error_rows=0)
        
        with pytest.raises(ValueError, match="revision_requires_correction"):
            await finalize_payment_validation(client, "revision_banco.xlsx", bank_code="banco_bogota", process_date=date(2026, 6, 25))

    asyncio.run(_run())

def test_finalize_bloquea_si_hoja_errores_tiene_filas():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        
        ctrl_path = resolve_bank_control_file_path("banco_bogota")
        client.downloaded_files[ctrl_path] = override_control_state("banco_bogota", "REVISION_CREADA")
        client.downloaded_files["revision/revision_banco.xlsx"] = create_finalize_excel_with_errors("banco_bogota", date(2026, 6, 25), error_rows=2)
        
        with pytest.raises(ValueError, match="revision_has_blocking_errors"):
            await finalize_payment_validation(client, "revision_banco.xlsx", bank_code="banco_bogota", process_date=date(2026, 6, 25))

    asyncio.run(_run())

def test_finalize_acepta_hoja_errores_vacia():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        
        ctrl_path = resolve_bank_control_file_path("banco_bogota")
        client.downloaded_files[ctrl_path] = override_control_state("banco_bogota", "REVISION_CREADA")
        client.downloaded_files["revision/revision_banco.xlsx"] = create_finalize_excel_with_errors("banco_bogota", date(2026, 6, 25), error_rows=0)
        
        with mock.patch("app.application.use_cases.payment_validation_finalize.register_payment_followups_after_finalize", new_callable=mock.AsyncMock):
            res = await finalize_payment_validation(client, "revision_banco.xlsx", bank_code="banco_bogota", process_date=date(2026, 6, 25))
        
        assert res["process_control_estado"] == "FINALIZADO"

    asyncio.run(_run())


def test_generate_no_retorna_already_generated_en_estado_correccion():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        ctrl_path = resolve_bank_control_file_path("banco_bogota")
        client.downloaded_files[ctrl_path] = override_control_state("banco_bogota", "REVISION_REQUIERE_CORRECCION")
        client.downloaded_files["revision/revision_banco.xlsx"] = create_finalize_excel_with_errors("banco_bogota", date(2026, 6, 25), error_rows=0)
        
        # Test 1: Mismo process_key y validation_file_path ya existen
        # Como tiene REVISION_REQUIERE_CORRECCION, NO DEBE abortar, sino continuar y fallar porque la carpeta no está vacía o intentará leer graph
        import httpx
        import zipfile
        try:
            await generate_payment_validation(client, date(2026, 6, 25), bank_code="banco_bogota")
            pytest.fail("Should have continued past early return and raised an error")
        except httpx.HTTPStatusError:
            pass # Prove it reached graph calls!
        except zipfile.BadZipFile:
            pass # Prove it reached openpyxl parsing of downloaded empty bank file!
        except ValueError as e:
            if str(e) == "review_folder_not_empty":
                pass # Or it reached folder empty check!
            else:
                raise

    asyncio.run(_run())

def test_generate_conserva_idempotencia_en_revision_creada():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        ctrl_path = resolve_bank_control_file_path("banco_bogota")
        client.downloaded_files[ctrl_path] = override_control_state("banco_bogota", "REVISION_CREADA")
        client.downloaded_files["revision/revision_banco.xlsx"] = create_finalize_excel_with_errors("banco_bogota", date(2026, 6, 25), error_rows=0)
        
        res = await generate_payment_validation(client, date(2026, 6, 25), bank_code="banco_bogota")
        
        assert res["already_generated"] is True
        assert res["file_action"] == "reused"

    asyncio.run(_run())

def test_generate_no_regenera_en_estados_avanzados():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        ctrl_path = resolve_bank_control_file_path("banco_bogota")
        client.downloaded_files[ctrl_path] = override_control_state("banco_bogota", "FINALIZADO")
        client.downloaded_files["revision/revision_banco.xlsx"] = create_finalize_excel_with_errors("banco_bogota", date(2026, 6, 25), error_rows=0)
        
        res = await generate_payment_validation(client, date(2026, 6, 25), bank_code="banco_bogota")
        
        assert res["already_generated"] is True
        assert res["file_action"] == "reused"

    asyncio.run(_run())
