"""Finalize: un ID Pago con varias filas de crédito (schema v2, monto banco único)."""

import asyncio
import io
from datetime import date

import openpyxl
import pytest

from app.application.services.review_schema import (
    AsientosPendientesCols,
    CasosPagoCols,
    DistribucionCols,
    EstadoPago,
    ReviewSheets,
    ValidarPago,
)
from app.application.use_cases.payment_validation_finalize import (
    SECRETARY_FIRST_DATA_ROW,
    SECRETARY_SHEET,
    finalize_payment_validation,
)
from tests.test_finalize_validation import (
    MockGraphClient,
    create_review_workbook,
    make_distrib_row,
    set_env_vars,
)


def _run(coro):
    return asyncio.run(coro)


def _monto_col_idx() -> int:
    return DistribucionCols.HEADERS.index(DistribucionCols.MONTO_BANCO)


def _make_group_row(
    credito: str,
    *,
    valor_int: float = 0,
    mora_a_aplicar: float = 0,
    abono_capital: float = 0,
    otros: float = 0,
    monto_banco=1000,
):
    row, hl = make_distrib_row(
        id_pago="TEST-001",
        cliente="CLI",
        credito=credito,
        monto_banco=monto_banco,
        valor_int=valor_int,
        mora_a_aplicar=mora_a_aplicar,
        abono_capital=abono_capital,
        otros=otros,
        estado=EstadoPago.NORMAL,
        validar_pago=ValidarPago.SI,
        extract_route=f"clientes/CLI/{credito}/extractos/e1.pdf",
    )
    return list(row), hl


def _make_success_workbook():
    mi = _monto_col_idx()
    r258, h258 = _make_group_row("258", valor_int=800, monto_banco=1000)
    r265, h265 = _make_group_row("265", mora_a_aplicar=50, monto_banco=None)
    r270, h270 = _make_group_row("270", abono_capital=150, monto_banco=None)
    assert r265[mi] is None
    assert r270[mi] is None
    casos = [["TEST-001", date(2026, 5, 10), "CLI", "concepto banco", 1000, ""]]
    return create_review_workbook(
        casos_data=casos,
        distrib_specs=[(r258, h258), (r265, h265), (r270, h270)],
    )


def _secretary_montos_and_total_validado(client: MockGraphClient) -> tuple[list[float | None], float]:
    sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
    wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
    ws = wb[SECRETARY_SHEET]
    col_monto = AsientosPendientesCols.HEADERS.index(AsientosPendientesCols.MONTO_BANCO) + 1
    col_total = AsientosPendientesCols.HEADERS.index(AsientosPendientesCols.TOTAL_VALIDADO) + 1
    montos = []
    total_validado = 0.0
    row = SECRETARY_FIRST_DATA_ROW
    while row <= ws.max_row:
        id_val = ws.cell(row, AsientosPendientesCols.HEADERS.index(AsientosPendientesCols.ID_PAGO) + 1).value
        if id_val in (None, "", "Total validado general"):
            break
        montos.append(ws.cell(row, col_monto).value)
        total_validado += float(ws.cell(row, col_total).value or 0)
        row += 1
    return montos, total_validado


def test_finalize_multi_row_payment_group_accepts_split_assignments():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.downloaded_files["revision/val_latest.xlsx"] = _make_success_workbook()
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"
        assert res["validated_payment_rows"] == 3

        payment_ids = set()
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        wb_hist = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = wb_hist[ReviewSheets.DISTRIBUCION]
        headers = [c.value for c in ws[1]]
        idx_id = headers.index(DistribucionCols.ID_PAGO)
        idx_cred = headers.index(DistribucionCols.CREDITO)
        creditos = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[idx_id]:
                continue
            payment_ids.add(str(row[idx_id]))
            creditos.append(str(row[idx_cred]))
        assert payment_ids == {"TEST-001"}
        assert sorted(creditos) == ["258", "265", "270"]

        grupos = len(payment_ids)
        assert grupos == 1

        montos_sec, total_validado = _secretary_montos_and_total_validado(client)
        assert len(montos_sec) == 3
        assert montos_sec[0] == pytest.approx(1000.0)
        assert montos_sec[1] in (None, "")
        assert montos_sec[2] in (None, "")
        assert total_validado == pytest.approx(1000.0)
        monto_bancario_total = sum(
            float(m) for m in montos_sec if m not in (None, "")
        )
        assert monto_bancario_total == pytest.approx(1000.0)

    _run(run_test())


def test_finalize_multi_row_payment_group_rejects_amount_shortfall():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        mi = _monto_col_idx()
        r258, h258 = _make_group_row("258", valor_int=800, monto_banco=1000)
        r265, h265 = _make_group_row("265", mora_a_aplicar=50, monto_banco=None)
        r270, h270 = _make_group_row("270", abono_capital=100, monto_banco=None)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            casos_data=[["TEST-001", date(2026, 5, 10), "CLI", "c", 1000, ""]],
            distrib_specs=[(r258, h258), (r265, h265), (r270, h270)],
        )
        with pytest.raises(ValueError, match="amount_mismatch"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())


def test_finalize_multi_row_payment_group_rejects_amount_excess():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r258, h258 = _make_group_row("258", valor_int=800, monto_banco=1000)
        r265, h265 = _make_group_row("265", mora_a_aplicar=100, monto_banco=None)
        r270, h270 = _make_group_row("270", abono_capital=150, monto_banco=None)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            casos_data=[["TEST-001", date(2026, 5, 10), "CLI", "c", 1000, ""]],
            distrib_specs=[(r258, h258), (r265, h265), (r270, h270)],
        )
        with pytest.raises(ValueError, match="amount_mismatch"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())


def test_finalize_rejects_duplicate_monto_banco_in_payment_group():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        mi = _monto_col_idx()
        r258, h258 = _make_group_row("258", valor_int=800, monto_banco=1000)
        r265, h265 = _make_group_row("265", mora_a_aplicar=50, monto_banco=None)
        r270, h270 = _make_group_row("270", abono_capital=150, monto_banco=None)
        r265[mi] = 1000
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            casos_data=[["TEST-001", date(2026, 5, 10), "CLI", "c", 1000, ""]],
            distrib_specs=[(r258, h258), (r265, h265), (r270, h270)],
        )
        with pytest.raises(ValueError, match="duplicate_bank_amount_in_payment_group"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())
