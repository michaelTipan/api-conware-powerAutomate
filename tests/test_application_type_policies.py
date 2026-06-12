"""Pruebas de políticas de Tipo Aplicación (4 valores visibles + legacy)."""

import asyncio
import io
from datetime import date
from unittest import mock

import openpyxl
import pytest

from app.application.services.review_schema import (
    APPLICATION_TYPES_SUPPORTED,
    ApplicationSubtype,
    CanonicalApplicationType,
    DistribucionAbonosCols,
    DistribucionCols,
    ExtractRole,
    ReviewSheets,
    TipoAplicacion,
    TipoAplicacionVisible,
    ValidarAbono,
    find_distribucion_pagos_sheet,
    normalize_tipo_aplicacion,
    parse_bank_tipo_aplicacion,
    resolve_application_policy,
)
from app.application.use_cases.payment_validation_generate import generate_payment_validation
from tests.test_generate_tipo_aplicacion import _run_with_pdf_mock
from tests.test_generate_validation import (
    MockGraphClient,
    STANDARD_BANK_HEADERS,
    create_bank_excel,
    create_excel,
    load_generated_workbook,
    set_env_vars,
    setup_client_structure,
    sheet_to_dicts,
)


@pytest.mark.parametrize("tipo", APPLICATION_TYPES_SUPPORTED)
def test_bank_accepts_visible_application_types(tipo):
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", tipo, "tx"]],
            default_tipo=tipo,
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            await generate_payment_validation(client, date(2026, 5, 26), bank_code="banco_bogota")
        assert client.uploaded_files

    asyncio.run(_run())


def test_bank_rejects_generic_abono():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 500000, "GEOEXCON", "ABONO", "tx"]],
            default_tipo="ABONO",
        )
        with pytest.raises(ValueError, match="generic_abono_not_supported"):
            with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
                await generate_payment_validation(client, date(2026, 5, 26), bank_code="banco_bogota")
        assert not client.uploaded_files

    asyncio.run(_run())


def test_legacy_abono_historical_readable():
    policy = resolve_application_policy("ABONO", from_bank=False)
    assert policy.legacy is True
    assert policy.subtipo_aplicacion == ApplicationSubtype.GENERAL
    assert normalize_tipo_aplicacion("ABONO") == TipoAplicacion.ABONO


def test_generate_creates_distribucion_pagos_not_legacy_distribucion():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "tx"]],
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            result = await generate_payment_validation(client, date(2026, 5, 26), bank_code="banco_bogota")
        wb = load_generated_workbook(client)
        assert ReviewSheets.DISTRIBUCION_PAGOS in wb.sheetnames
        assert ReviewSheets.DISTRIBUCION not in wb.sheetnames
        assert result["distribution_payments_sheet"] == ReviewSheets.DISTRIBUCION_PAGOS

    asyncio.run(_run())


def test_generate_creates_distribucion_abonos():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 500000, "GEOEXCON", "ABONO CAPITAL", "tx"]],
            default_tipo="ABONO CAPITAL",
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            await generate_payment_validation(client, date(2026, 5, 26), bank_code="banco_bogota")
        wb = load_generated_workbook(client)
        assert ReviewSheets.DISTRIBUCION_ABONOS in wb.sheetnames

    asyncio.run(_run())


def test_finalize_legacy_distribucion_fallback():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.DISTRIBUCION_LEGACY
    ws.append(list(DistribucionCols.HEADERS))
    found = find_distribucion_pagos_sheet(wb)
    assert found.title == ReviewSheets.DISTRIBUCION_LEGACY


@pytest.mark.parametrize(
    "tipo,expected_sheet",
    [
        ("PAGO", "pagos"),
        ("PAGO Y ABONO CAPITAL", "pagos"),
        ("ABONO CAPITAL", "abonos"),
        ("ABONO MORA", "abonos"),
    ],
)
def test_generate_routes_by_application_type(tipo, expected_sheet):
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_bank_excel(
            [[date(2025, 12, 23), 25443565, "GEOEXCON", tipo, "tx"]],
            default_tipo=tipo,
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            await generate_payment_validation(client, date(2026, 5, 26), bank_code="banco_bogota")
        wb = load_generated_workbook(client)
        pagos = sheet_to_dicts(wb[ReviewSheets.DISTRIBUCION_PAGOS])
        abonos = sheet_to_dicts(wb[ReviewSheets.DISTRIBUCION_ABONOS])
        if expected_sheet == "pagos":
            assert pagos
            assert not any(r.get(DistribucionAbonosCols.ID_PAGO) for r in abonos if abonos)
        else:
            assert abonos
            assert all(r.get(DistribucionAbonosCols.TIPO_APLICACION_ORIGINAL) == tipo for r in abonos)
            assert not pagos

    asyncio.run(_run())


def test_abono_capital_no_extract_required_in_policy():
    p = parse_bank_tipo_aplicacion("ABONO CAPITAL")
    assert p.requiere_extracto is False
    assert p.rol_extracto == ExtractRole.NO_APLICA


def test_abono_mora_requires_reference_extract_policy():
    p = parse_bank_tipo_aplicacion("ABONO MORA")
    assert p.requiere_extracto is True
    assert p.rol_extracto == ExtractRole.REFERENCIA_MORA
    assert p.actualiza_ibr is False


def test_pago_y_abono_capital_policy():
    p = parse_bank_tipo_aplicacion("PAGO Y ABONO CAPITAL")
    assert p.tipo_aplicacion_canonica == CanonicalApplicationType.PAGO
    assert p.subtipo_aplicacion == ApplicationSubtype.CUOTA_MAS_CAPITAL
    assert p.actualiza_ibr is True


def test_pago_atrasado_still_closes_cuota_policy():
    p = resolve_application_policy("PAGO", from_bank=False)
    assert p.cierra_cuota is True
    assert p.actualiza_ibr is True


def test_abono_mora_does_not_close_cuota_policy():
    p = resolve_application_policy("ABONO MORA", from_bank=False)
    assert p.cierra_cuota is False
    assert p.genera_siguiente_extracto is False


def test_generate_response_lists_supported_types():
    async def _run():
        set_env_vars()
        client = MockGraphClient()
        client.children = []
        setup_client_structure(client)
        client.downloaded_files["banco.xlsx"] = create_excel(
            STANDARD_BANK_HEADERS,
            [
                [date(2025, 12, 23), 25443565, "GEOEXCON", "PAGO", "tx-p"],
                [date(2025, 12, 24), 100000, "GEOEXCON", "ABONO CAPITAL", "tx-a"],
            ],
        )
        with _run_with_pdf_mock(client)[0], _run_with_pdf_mock(client)[1]:
            result = await generate_payment_validation(client, date(2026, 5, 26), bank_code="banco_bogota")
        assert result["application_types_supported"] == list(APPLICATION_TYPES_SUPPORTED)
        assert "Distribucion_Pagos" in result.get("user_message", "")
        assert "Distribucion_Abonos" in result.get("user_message", "")

    asyncio.run(_run())
