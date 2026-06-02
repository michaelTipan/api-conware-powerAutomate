"""
Lectura/actualización del control oficial por banco (fila 2, hoja Procesos).

Única fuente operativa del flujo Generate → Apply.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any

from openpyxl import load_workbook

from app.application.sharepoint_resolution import encode_graph_drive_path
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    PROCESS_CONTROL_COLUMNS,
    SHEET_NAME,
)
from app.domain.ports.graph import GraphApiPort


BANK_CODE_BOGOTA = "banco_bogota"
BANK_CODE_BANCOLOMBIA = "banco_bancolombia"


def normalize_bank_code(bank_code: str | None) -> str:
    bc = (bank_code or "").strip()
    return bc or BANK_CODE_BOGOTA


def validate_bank_code(bank_code: str) -> None:
    if bank_code not in (BANK_CODE_BOGOTA, BANK_CODE_BANCOLOMBIA):
        raise ValueError("invalid_bank_code")


def resolve_process_control_path_for_bank(bank_code: str) -> str:
    if bank_code == BANK_CODE_BOGOTA:
        return PROCESS_CONTROL_BANK_FILE_BOGOTA
    if bank_code == BANK_CODE_BANCOLOMBIA:
        return PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA
    raise ValueError("invalid_bank_code")


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _content_endpoint(site_id: str, drive_id: str, rel_path: str) -> str:
    enc = encode_graph_drive_path(rel_path.strip().strip("/"))
    return f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/content"


def _col_index(name: str) -> int:
    return PROCESS_CONTROL_COLUMNS.index(name) + 1


def _is_active_cell(value: Any) -> bool:
    if value is True:
        return True
    if value in (False, None):
        return False
    s = str(value).strip().lower()
    return s in ("true", "1", "yes", "si", "sí", "x")


@dataclass(frozen=True)
class ProcessControlSnapshot:
    control_file_path: str
    estado_proceso: str
    is_active: bool
    process_key: str
    validation_file_path: str
    historical_file_path: str
    secretary_file_path: str
    email_pdf_path: str
    notify_idempotency_key: str
    merge_manifest_path: str
    merge_idempotency_key: str
    apply_idempotency_key: str
    bank_code: str
    bank_name: str


def parse_process_control_row2(raw: bytes, *, control_file_path: str) -> ProcessControlSnapshot:
    wb = load_workbook(filename=io.BytesIO(raw), data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError("process_control_invalid_structure")
        ws = wb[SHEET_NAME]
        if ws.max_row < 2:
            raise ValueError("process_control_invalid_structure")

        estado = str(ws.cell(row=2, column=_col_index("EstadoProceso")).value or "").strip()
        is_active = _is_active_cell(ws.cell(row=2, column=_col_index("IsActive")).value)
        pkey = str(ws.cell(row=2, column=_col_index("ProcessKey")).value or "").strip()
        vpath = str(ws.cell(row=2, column=_col_index("ValidationFilePath")).value or "").strip().strip("/")
        hpath = str(ws.cell(row=2, column=_col_index("HistoricalFilePath")).value or "").strip().strip("/")
        spath = str(ws.cell(row=2, column=_col_index("SecretaryFilePath")).value or "").strip().strip("/")
        epdf = str(ws.cell(row=2, column=_col_index("EmailPdfPath")).value or "").strip().strip("/")
        nid = str(ws.cell(row=2, column=_col_index("NotifyIdempotencyKey")).value or "").strip()
        mmp = str(ws.cell(row=2, column=_col_index("MergeManifestPath")).value or "").strip().strip("/")
        mid = str(ws.cell(row=2, column=_col_index("MergeIdempotencyKey")).value or "").strip()
        aid = str(ws.cell(row=2, column=_col_index("ApplyIdempotencyKey")).value or "").strip()
        bc = str(ws.cell(row=2, column=_col_index("BankCode")).value or "").strip()
        bn = str(ws.cell(row=2, column=_col_index("BankName")).value or "").strip()

        return ProcessControlSnapshot(
            control_file_path=control_file_path,
            estado_proceso=estado,
            is_active=is_active,
            process_key=pkey,
            validation_file_path=vpath,
            historical_file_path=hpath,
            secretary_file_path=spath,
            email_pdf_path=epdf,
            notify_idempotency_key=nid,
            merge_manifest_path=mmp,
            merge_idempotency_key=mid,
            apply_idempotency_key=aid,
            bank_code=bc,
            bank_name=bn,
        )
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


async def download_process_control_bytes(
    graph: GraphApiPort, site_id: str, drive_id: str, *, bank_code: str
) -> tuple[str, bytes]:
    rel = resolve_process_control_path_for_bank(bank_code).strip().strip("/")
    raw = await graph.get_bytes(_content_endpoint(site_id, drive_id, rel))
    return rel, raw


async def read_process_control_snapshot(
    graph: GraphApiPort, site_id: str, drive_id: str, *, bank_code: str
) -> ProcessControlSnapshot:
    rel, raw = await download_process_control_bytes(graph, site_id, drive_id, bank_code=bank_code)
    return parse_process_control_row2(raw, control_file_path=rel)


async def update_process_control_row2(
    graph: GraphApiPort,
    site_id: str,
    drive_id: str,
    *,
    bank_code: str,
    updates: dict[str, Any],
) -> None:
    rel = resolve_process_control_path_for_bank(bank_code).strip().strip("/")
    raw = await graph.get_bytes(_content_endpoint(site_id, drive_id, rel))
    wb = load_workbook(filename=io.BytesIO(raw), data_only=False)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError("process_control_invalid_structure")
        ws = wb[SHEET_NAME]
        if ws.max_row < 2:
            raise ValueError("process_control_invalid_structure")

        for key, val in updates.items():
            if key not in PROCESS_CONTROL_COLUMNS:
                continue
            ws.cell(row=2, column=_col_index(key), value=val)

        buf = io.BytesIO()
        wb.save(buf)
        out = buf.getvalue()
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()

    await graph.put_bytes(
        _content_endpoint(site_id, drive_id, rel),
        out,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

