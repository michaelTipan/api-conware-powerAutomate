"""Fórmulas O:P y protección de _AUTOMATION_LOG en tablas de amortización."""

from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from app.application.services.amortization_workbook import (
    AUTOMATION_LOG_SHEET,
    append_automation_log,
    ensure_application_related_formulas,
    find_header_row,
    protect_automation_log_sheet,
    ensure_automation_log,
    _is_formula_value,
)

_AMORT_HEADER_ROW = [
    "dia",
    "mes",
    "año",
    "IBR +i",
    "Fecha pago",
    "Valor intereses",
    "Abono a K",
    "intereses mora",
    "Valor pagado cliente",
    "Saldos Menores",
]


def _ws_with_op_formulas(
    *,
    formula_through_row: int = 7,
    p_multipliers: dict[int, int] | None = None,
    max_row: int = 10,
) -> openpyxl.worksheet.worksheet.Worksheet:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EQUINORTE"
    ws.append(_AMORT_HEADER_ROW)
    while (ws.max_row or 1) < max_row:
        ws.append([None] * len(_AMORT_HEADER_ROW))
    for r in range(4, formula_through_row + 1):
        ws.cell(r, 3, 30.0)
        ws.cell(r, 15, f"=+C{r}/30")
        mult = (p_multipliers or {}).get(r, 8 if r < 7 else 10)
        ws.cell(r, 16, f"=+O{r}*{mult}")
    return ws


def test_ensure_op_formulas_extends_to_max_application_row():
    ws = _ws_with_op_formulas(formula_through_row=7, max_row=9)
    header_row = find_header_row(ws)
    result = ensure_application_related_formulas(
        ws, max_application_row=9, header_row=header_row
    )
    assert result.last_row == 9
    assert result.rows_filled >= 1
    assert _is_formula_value(ws.cell(9, 15).value)
    assert _is_formula_value(ws.cell(9, 16).value)
    assert str(ws.cell(9, 15).value).upper() == "=+C9/30"
    assert str(ws.cell(9, 16).value).upper() == "=+O9*10"


def test_ensure_op_formulas_row_10_from_row_9_template():
    ws = _ws_with_op_formulas(formula_through_row=9, max_row=10)
    header_row = find_header_row(ws)
    ensure_application_related_formulas(ws, max_application_row=10, header_row=header_row)
    assert str(ws.cell(10, 15).value).upper() == "=+C10/30"
    assert str(ws.cell(10, 16).value).upper() == "=+O10*10"


def test_ensure_op_formulas_does_not_overwrite_existing():
    ws = _ws_with_op_formulas(formula_through_row=8, max_row=9)
    ws.cell(9, 15, "=+CUSTOM9")
    ws.cell(9, 16, "=+CUSTOMP9")
    header_row = find_header_row(ws)
    ensure_application_related_formulas(ws, max_application_row=9, header_row=header_row)
    assert ws.cell(9, 15).value == "=+CUSTOM9"
    assert ws.cell(9, 16).value == "=+CUSTOMP9"


def test_ensure_op_formulas_no_op_when_max_row_at_header():
    ws = _ws_with_op_formulas(formula_through_row=7, max_row=8)
    header_row = find_header_row(ws)
    before_o8 = ws.cell(8, 15).value
    result = ensure_application_related_formulas(
        ws, max_application_row=header_row, header_row=header_row
    )
    assert result.cells_filled == 0
    assert ws.cell(8, 15).value == before_o8


def test_automation_log_created_and_protected():
    wb = openpyxl.Workbook()
    ensure_automation_log(wb)
    protected, warning = protect_automation_log_sheet(wb)
    assert protected is True
    assert warning is None
    ws = wb[AUTOMATION_LOG_SHEET]
    assert ws.protection.sheet is True
    assert ws.cell(1, 1).protection.locked is True


def test_append_automation_log_on_protected_sheet(monkeypatch):
    monkeypatch.delenv("AMORTIZATION_LOG_SHEET_PROTECTION_PASSWORD", raising=False)
    wb = openpyxl.Workbook()
    ensure_automation_log(wb)
    protect_automation_log_sheet(wb)
    append_automation_log(
        wb,
        {
            "id_pago": "p1",
            "cliente": "c",
            "credito": "cr",
            "application_row": 8,
            "accion": "APLICADO",
            "idempotency_key": "k1",
        },
    )
    ws = wb[AUTOMATION_LOG_SHEET]
    assert ws.max_row >= 2
    protect_automation_log_sheet(wb)
    assert ws.protection.sheet is True


def test_automation_log_protection_password_from_env(monkeypatch):
    monkeypatch.setenv("AMORTIZATION_LOG_SHEET_PROTECTION_PASSWORD", "test-secret")
    wb = openpyxl.Workbook()
    ensure_automation_log(wb)
    protect_automation_log_sheet(wb)
    ws = wb[AUTOMATION_LOG_SHEET]
    assert ws.protection.sheet is True
    assert bool(ws.protection.password)
