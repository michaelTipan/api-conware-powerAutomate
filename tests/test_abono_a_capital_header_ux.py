"""UX: encabezado «Abono a capital» en workbooks nuevos; lectura legacy «Mora a aplicar»."""

import asyncio
import io
from datetime import date, datetime

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.application.services.review_schema import (
    ApplicationSubtype,
    DistribucionCols,
    ReviewSheets,
    TipoAplicacionVisible,
    normalize_distrib_row_keys,
    resolve_application_policy,
)
from app.application.use_cases.payment_validation_finalize import finalize_payment_validation
from tests.test_finalize_validation import (
    MockGraphClient,
    create_review_workbook,
    make_distrib_row,
    set_env_vars,
)
from tests.test_generate_validation import (
    _dist_col,
    _first_data_row,
    _header_row_index,
    run_generate,
)
from tests.test_pago_y_abono_capital_finalize import (
    make_pago_y_abono_capital_row,
)


def _run(coro):
    return asyncio.run(coro)


def _legacy_distrib_headers() -> list[str]:
    headers = list(DistribucionCols.HEADERS)
    idx = headers.index(DistribucionCols.ABONO_A_CAPITAL)
    headers[idx] = DistribucionCols.MORA_A_APLICAR
    return headers


def test_generate_new_workbook_uses_abono_a_capital_header():
    _, _, wb = run_generate(
        [[datetime(2025, 12, 23), 25443565, "GEOEXCON", ""]],
        date(2025, 12, 23),
    )
    ws = wb[ReviewSheets.DISTRIBUCION_PAGOS]
    hr = _header_row_index(ws, DistribucionCols.ID_PAGO)
    headers = [c.value for c in ws[hr]]
    assert DistribucionCols.ABONO_A_CAPITAL in headers
    assert DistribucionCols.MORA_A_APLICAR not in headers


def test_generate_does_not_emit_both_capital_headers():
    _, _, wb = run_generate(
        [[datetime(2025, 12, 23), 25443565, "GEOEXCON", ""]],
        date(2025, 12, 23),
    )
    ws = wb[ReviewSheets.DISTRIBUCION_PAGOS]
    hr = _header_row_index(ws, DistribucionCols.ID_PAGO)
    headers = [str(c.value or "") for c in ws[hr]]
    assert headers.count(DistribucionCols.ABONO_A_CAPITAL) == 1
    assert DistribucionCols.MORA_A_APLICAR not in headers


def test_saldo_formula_references_abono_a_capital_column():
    _, _, wb = run_generate(
        [[datetime(2025, 12, 23), 25443565, "GEOEXCON", ""]],
        date(2025, 12, 23),
    )
    ws = wb[ReviewSheets.DISTRIBUCION_PAGOS]
    dr = _first_data_row(ws, DistribucionCols.ID_PAGO)
    formula = ws.cell(dr, _dist_col(DistribucionCols.SALDO_POR_ASIGNAR)).value
    col_j = get_column_letter(_dist_col(DistribucionCols.ABONO_A_CAPITAL))
    assert isinstance(formula, str) and formula.startswith("=")
    assert f"${col_j}$" in formula


def test_finalize_accepts_new_workbook_with_abono_a_capital_header():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_pago_y_abono_capital_row(valor_extracto=800, abono_capital=200)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)]
        )
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"

    _run(run_test())


def test_finalize_accepts_legacy_workbook_with_mora_a_aplicar_header():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(valor_int=800, abono_k=200, mora=0, monto_banco=1000)
        wb = openpyxl.Workbook()
        ws_ctrl = wb.active
        ws_ctrl.title = ReviewSheets.CONTROL
        ws_ctrl.append(["Campo", "Valor"])
        ws_ctrl.append(["Procesar", "SI"])
        ws_ctrl.append(["Estado", "EN_REVISION"])
        ws_casos = wb.create_sheet(ReviewSheets.CASOS_PAGO)
        ws_casos.append(["ID Pago", "Fecha banco", "Cliente", "Concepto banco", "Monto banco", "Observación"])
        ws_casos.append(["ID1", None, "CLI", "c", 1000, ""])
        ws_dist = wb.create_sheet(ReviewSheets.DISTRIBUCION)
        legacy_headers = _legacy_distrib_headers()
        ws_dist.append(legacy_headers)
        vals = dict(zip(DistribucionCols.HEADERS, r))
        legacy_row = [
            vals.get(DistribucionCols.ABONO_A_CAPITAL, "")
            if h == DistribucionCols.MORA_A_APLICAR
            else vals.get(h, "")
            for h in legacy_headers
        ]
        ws_dist.append(legacy_row)
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"

    _run(run_test())


def test_normalize_distrib_row_keys_maps_legacy_mora_header():
    row = normalize_distrib_row_keys(
        {
            DistribucionCols.MORA_A_APLICAR: 200,
            DistribucionCols.APLICAR_A_EXTRACTO: 800,
        }
    )
    assert row[DistribucionCols.ABONO_A_CAPITAL] == 200
    assert DistribucionCols.ABONO_A_CAPITAL in row


def test_abono_mora_policy_does_not_use_capital_column_for_subtype():
    policy = resolve_application_policy(TipoAplicacionVisible.ABONO_MORA, from_bank=False)
    assert policy.subtipo_aplicacion == ApplicationSubtype.MORA
    assert policy.tipo_aplicacion_canonica == "ABONO"


def test_legacy_historical_row_keys_remain_readable():
    legacy = {
        DistribucionCols.MORA_A_APLICAR: 50,
        DistribucionCols.APLICAR_A_EXTRACTO: 100,
        DistribucionCols.OTROS_VALORES: 10,
    }
    normalized = normalize_distrib_row_keys(legacy)
    assert normalized[DistribucionCols.ABONO_A_CAPITAL] == 50
