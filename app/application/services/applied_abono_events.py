"""
Identificación de eventos ABONO ya aplicados vía ``_AUTOMATION_LOG`` (retry parcial).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.application.services.amortization_workbook import (
    ADOPTADO_EXISTENTE,
    APLICADO,
    AUTOMATION_LOG_SHEET,
    _automation_log_header_map,
    _cell_float,
    detect_amortization_sheet,
    detect_headers,
    find_header_row,
)
from app.application.services.review_schema import TipoAplicacion, normalize_credito_digits
from app.application.use_cases.merge_composite_validado_pdfs import normalize_sharepoint_path

APPLIED_ABONO_AMOUNT_UNRESOLVABLE = "APPLIED_ABONO_AMOUNT_UNRESOLVABLE"
APPLIED_ABONO_EVENT_AMBIGUOUS = "APPLIED_ABONO_EVENT_AMBIGUOUS"

_APPLIED_ACCIONES = frozenset({APLICADO, ADOPTADO_EXISTENTE})
_FAILED_ACCIONES = frozenset({"ERROR", "PENDIENTE", "REVISION_MANUAL"})


@dataclass(frozen=True)
class AppliedAbonoEventSnapshot:
    idempotency_key: str
    id_pago: str
    credito: str
    asiento_pdf_path: str
    asiento_pdf_hash: str
    application_row: int | None
    accion: str
    tipo_aplicacion: str
    valor_pagado_cliente: Decimal | None
    log_row: int


def _norm_credito_token(credito: str) -> str:
    import re

    digits = normalize_credito_digits(credito) or re.sub(r"\D", "", str(credito or ""))
    return digits or str(credito or "").strip()


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if not text:
        return None
    try:
        return Decimal(text.replace(",", ""))
    except InvalidOperation:
        num = _cell_float(value)
        if num is None:
            return None
        return Decimal(str(num))


def _is_applied_snapshot(snap: AppliedAbonoEventSnapshot) -> bool:
    accion = str(snap.accion or "").strip().upper()
    estado_failed = accion in _FAILED_ACCIONES
    if estado_failed:
        return False
    if accion in _APPLIED_ACCIONES:
        return True
    estado = accion
    return estado in _APPLIED_ACCIONES


def _snapshot_tipo_is_abono(tipo: str) -> bool:
    raw = str(tipo or "").strip().upper()
    return not raw or raw == TipoAplicacion.ABONO.value


def load_applied_abono_event_snapshots(
    workbook: Workbook,
    *,
    id_pago_filter: str | None = None,
) -> list[AppliedAbonoEventSnapshot]:
    """Lee eventos ABONO aplicados desde ``_AUTOMATION_LOG``."""
    if AUTOMATION_LOG_SHEET not in workbook.sheetnames:
        return []
    ws = workbook[AUTOMATION_LOG_SHEET]
    if (ws.max_row or 1) < 2:
        return []
    header_map = _automation_log_header_map(ws)
    if not header_map.get("IdempotencyKey"):
        return []

    id_filter = str(id_pago_filter or "").strip()
    snapshots: list[AppliedAbonoEventSnapshot] = []

    def _cell(row: int, name: str) -> str:
        col = header_map.get(name)
        if col is None:
            return ""
        return str(ws.cell(row, col).value or "").strip()

    for row in range(2, (ws.max_row or 1) + 1):
        key = _cell(row, "IdempotencyKey")
        if not key:
            continue
        id_pago = _cell(row, "IdPago")
        if id_filter and id_pago != id_filter:
            continue
        tipo = _cell(row, "TipoAplicacion")
        if not _snapshot_tipo_is_abono(tipo):
            continue
        accion = _cell(row, "Accion") or _cell(row, "Estado")
        app_raw = _cell(row, "ApplicationRow") or _cell(row, "Fila")
        application_row = int(app_raw) if app_raw.isdigit() else None
        vp_raw = _cell(row, "ValorPagadoCliente")
        valor = _parse_decimal(vp_raw) if vp_raw else None
        snapshots.append(
            AppliedAbonoEventSnapshot(
                idempotency_key=key,
                id_pago=id_pago,
                credito=_cell(row, "Credito"),
                asiento_pdf_path=_cell(row, "AsientoPdfPath"),
                asiento_pdf_hash=_cell(row, "AsientoPdfHash"),
                application_row=application_row,
                accion=accion,
                tipo_aplicacion=tipo,
                valor_pagado_cliente=valor,
                log_row=row,
            )
        )
    return [s for s in snapshots if _is_applied_snapshot(s)]


def read_valor_pagado_cliente_from_application_row(
    ws: Worksheet,
    row: int,
    headers: dict[str, int],
) -> Decimal | None:
    """Lee el monto canónico ya escrito en la fila de Aplicación de Pagos."""
    col = headers.get("valor_pagado_cliente")
    if col is None:
        return None
    raw = ws.cell(row, col).value
    if raw is None or str(raw).strip() == "":
        return None
    if isinstance(raw, str) and raw.strip().startswith("="):
        return None
    num = _cell_float(raw)
    if num is None or abs(num) <= 0:
        return None
    return Decimal(str(num))


def resolve_applied_abono_amount(
    snapshot: AppliedAbonoEventSnapshot,
    *,
    ws: Worksheet | None = None,
    headers: dict[str, int] | None = None,
) -> tuple[Decimal | None, str | None]:
    """
    Monto del evento aplicado: log ``ValorPagadoCliente`` o fila de aplicación.
    """
    if snapshot.valor_pagado_cliente is not None and snapshot.valor_pagado_cliente > 0:
        return snapshot.valor_pagado_cliente, None
    app_row = snapshot.application_row
    if ws is not None and headers is not None and app_row is not None:
        amount = read_valor_pagado_cliente_from_application_row(ws, app_row, headers)
        if amount is not None and amount > 0:
            return amount, None
    return None, APPLIED_ABONO_AMOUNT_UNRESOLVABLE


def _paths_match(a: str, b: str) -> bool:
    na = normalize_sharepoint_path(a)
    nb = normalize_sharepoint_path(b)
    return bool(na and nb and na == nb)


def _credito_matches(snap_cred: str, target: str) -> bool:
    a = _norm_credito_token(snap_cred)
    b = _norm_credito_token(target)
    if a and b and a == b:
        return True
    return str(snap_cred or "").strip() == str(target or "").strip()


def find_applied_abono_event(
    snapshots: list[AppliedAbonoEventSnapshot],
    *,
    id_pago: str,
    credito: str,
    asiento_pdf_path: str,
    idempotency_key: str | None = None,
    asiento_pdf_hash: str | None = None,
) -> tuple[AppliedAbonoEventSnapshot | None, str | None]:
    """
    Localiza un evento ya aplicado. Devuelve (match, error_code) si hay ambigüedad.
    """
    id_pago = str(id_pago or "").strip()
    credito_norm = _norm_credito_token(credito) or str(credito or "").strip()
    path_norm = normalize_sharepoint_path(asiento_pdf_path)
    key = str(idempotency_key or "").strip()
    pdf_hash = str(asiento_pdf_hash or "").strip()

    candidates: list[AppliedAbonoEventSnapshot] = []

    if key:
        for snap in snapshots:
            if snap.idempotency_key == key:
                if pdf_hash and snap.asiento_pdf_hash and snap.asiento_pdf_hash != pdf_hash:
                    continue
                candidates.append(snap)
    else:
        for snap in snapshots:
            if id_pago and snap.id_pago and snap.id_pago != id_pago:
                continue
            if not _credito_matches(snap.credito, credito_norm):
                continue
            if not _paths_match(snap.asiento_pdf_path, path_norm):
                continue
            if pdf_hash and snap.asiento_pdf_hash and snap.asiento_pdf_hash != pdf_hash:
                continue
            candidates.append(snap)

    if not candidates:
        return None, None
    if len(candidates) > 1:
        return None, APPLIED_ABONO_EVENT_AMBIGUOUS
    return candidates[0], None


def load_applied_abono_snapshots_from_table_bytes(
    tabla_bytes: bytes,
    *,
    id_pago: str,
    tabla_path: str = "",
) -> tuple[list[AppliedAbonoEventSnapshot], Worksheet | None, dict[str, int], int]:
    """Abre tabla y devuelve snapshots + hoja amortización para lectura de montos."""
    wb = openpyxl.load_workbook(io.BytesIO(tabla_bytes), data_only=True)
    try:
        match = detect_amortization_sheet(wb, tabla_amortizacion_path=tabla_path)
        ws = match.worksheet
        headers = match.headers
        header_row = match.header_row
        snapshots = load_applied_abono_event_snapshots(wb, id_pago_filter=id_pago)
        return snapshots, ws, headers, header_row
    except Exception:
        header_row = find_header_row(wb.active)
        headers = detect_headers(wb.active, header_row=header_row)
        snapshots = load_applied_abono_event_snapshots(wb, id_pago_filter=id_pago)
        return snapshots, wb.active, headers, header_row
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()
