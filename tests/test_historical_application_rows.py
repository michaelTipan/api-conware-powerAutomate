"""Tests lectura compartida PAGO/ABONO en histórico."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.application.services.historical_application_rows import (
    abono_row_included_for_notify_merge,
    detect_application_type_group_conflict,
    group_abono_rows_for_email,
    group_rows_by_id_pago,
    read_validated_abono_rows,
    read_validated_payment_rows,
)
from app.application.services.review_schema import (
    DistribucionAbonosCols,
    DistribucionCols,
    ReviewSheets,
    TipoAplicacion,
    ValidarAbono,
    ValidarPago,
)
from tests.test_finalize_validation import make_abono_row, make_distrib_row


def _payment_historical_bytes(*, include_abonos: bool = False) -> bytes:
    r, _ = make_distrib_row(id_pago="P1", validar_pago=ValidarPago.SI, ruta_pdf_internal="ext/p1.pdf")
    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.DISTRIBUCION
    ws.append(DistribucionCols.HEADERS)
    ws.append(r)
    if include_abonos:
        ws_ab = wb.create_sheet(ReviewSheets.DISTRIBUCION_ABONOS)
        ws_ab.append(
            list(DistribucionAbonosCols.HEADERS) + [DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES]
        )
        a1, _ = make_abono_row(id_pago="AB1", validar_abono=ValidarAbono.SI, credito="258")
        a2, _ = make_abono_row(id_pago="AB1", validar_abono=ValidarAbono.SI, credito="265")
        for row in (a1, a2):
            row = list(row) + [f"clientes/CLI/CREDITO# {row[2]}/ASIENTOS CONTABLES"]
            ws_ab.append(row)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_read_payment_rows_legacy_distribucion_only():
    wb = load_workbook(BytesIO(_payment_historical_bytes()), data_only=True)
    rows = read_validated_payment_rows(wb, legacy_estado_token="VALIDAR")
    assert len(rows) == 1
    assert rows[0]["tipo_aplicacion"] == TipoAplicacion.PAGO.value
    assert rows[0]["requiere_extracto"] is True


def test_read_abono_rows_empty_when_sheet_missing():
    wb = load_workbook(BytesIO(_payment_historical_bytes(include_abonos=False)), data_only=True)
    assert read_validated_abono_rows(wb) == []


def test_group_abono_single_row_per_id_pago():
    wb = load_workbook(BytesIO(_payment_historical_bytes(include_abonos=True)), data_only=True)
    abono_rows = read_validated_abono_rows(wb)
    assert len(abono_rows) == 2
    groups = group_abono_rows_for_email(abono_rows)
    assert len(groups) == 1
    assert groups[0].id_pago == "AB1"
    assert groups[0].creditos_seleccionados == ("258", "265")


def test_detect_application_type_group_conflict():
    payment_groups = {"X1": [{"id_pago": "X1"}]}
    abono_groups = {"X1": [{"id_pago": "X1"}]}
    with pytest.raises(ValueError, match="application_type_group_conflict"):
        detect_application_type_group_conflict(payment_groups, abono_groups)


def test_abono_inclusion_requires_ruta_asientos():
    wb = Workbook()
    ws = wb.create_sheet(ReviewSheets.DISTRIBUCION_ABONOS)
    ws.append(list(DistribucionAbonosCols.HEADERS) + [DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES])
    a_no, _ = make_abono_row(validar_abono=ValidarAbono.SI)
    ws.append(list(a_no) + [""])
    bio = BytesIO()
    wb.save(bio)
    wb2 = load_workbook(BytesIO(bio.getvalue()), data_only=True)
    assert read_validated_abono_rows(wb2) == []
